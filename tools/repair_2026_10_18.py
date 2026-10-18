# -*- coding: utf-8 -*-
"""Morning fix 2026-10-18. ROWS ARE REMOVED from Recalls — the commit
message carries "remove-rows".

Rebuilt by the operator's review session from the morning run's write-up
(Drive doc "morning-fix 2026-10-18"); every fact re-checked on main.

A. FIVE ROWS OUT OF RECALLS, each archived to Rejected (append-only):
   * NAFDAC public-alert-no-023-2025 — multivitamin supplements; released
     5 July 2025; missing child-resistant packaging (packaging, out of
     scope). Removed 2026-10-17, RESURRECTED by the Recalls row-merge.
   * NAFDAC public-alert-no-027-2022 — Japata alcoholic bitters; released
     10 May 2022; no in-scope hazard named. Removed 2026-10-17, RESURRECTED.
   * NAFDAC public-alert-no-041-2024 — Nivea roll-on DEODORANT; released
     30 Oct 2024, category Cosmetics, relayed from the EU Safety Gate for
     dangerous NON-FOOD products. Not food (R1). repair_2026_10_17 stamped
     its archive entry SUPERSEDED in favour of the published copy; that
     stamp is withdrawn here.
   * CFS (HK) whatsnew_fa/2026_633 — duplicate of press/20261012_12665
     (preserved bean curd, Bacillus cereus 140,000/g). RESURRECTED.
   * CFS (HK) whatsnew_fa/2026_632 — duplicate of press/20261009_12660
     (Marks & Spencer salad and pesto, Salmonella). RESURRECTED.
   The resurrection path is closed in pipeline/xlsx_merge.py (same upload).

B. BILINGUAL RASFF SUBJECTS. Ten published RASFF rows carry the regulator's
   bilingual subject whole in Product and/or Reason. split_bilingual() keeps
   RASFF's own English half; nothing is translated; each original goes to
   Notes. The writer guard that should have done this was unreachable
   (pipeline/merge_master.py, same upload).

C. RappelConso fiche 23785 (ref. 2026-10-0114), held in Pending on an empty
   Company with Product "cr20cz" (a scraper artifact). Read on its own page
   2026-10-18: published 16/10/2026; dénomination "Crevette entière cuite
   réfrigérée 20/30 sauvage penaeus indicus péchée en océan indien ouest
   FAO51"; marque "Sans marque"; lot 9262800232; DLC 16/10/2026; sold
   07/10–13/10/2026; origine de la fiche "CRUSTA C" (the same field the
   published sibling fiche 23786 took its Company from); motif "Présence de
   listeria monocytogenes".
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
TODAY = "2026-10-18"
OP = "[morning fix 2026-10-18]"

REMOVE = {
    "public-alert-no-023-2025": (
        "not in scope: NAFDAC public alert 023/2025, released 5 July 2025 — missing "
        "child-resistant packaging on iron-containing multivitamins (packaging). Dated "
        "2026 only because that is when it was found. Removed 2026-10-17 and resurrected "
        "by the Recalls row-merge; removed again. " + OP),
    "public-alert-no-027-2022": (
        "not in scope: NAFDAC public alert 027/2022, released 10 May 2022 — 'adulterated, "
        "unwholesome' alcoholic bitters, no in-scope hazard named. Removed 2026-10-17 and "
        "resurrected by the Recalls row-merge; removed again. " + OP),
    "public-alert-no-041-2024": (
        "not_food: NAFDAC public alert 041/2024, released 30 Oct 2024, category Cosmetics — "
        "Nivea roll-on deodorant (BMHCA), relayed from the EU Safety Gate for dangerous "
        "non-food products. Not food (R1). " + OP),
    "whatsnew_fa/2026_633.html": (
        "duplicate: CFS (HK) Food Alert 2026_633 is the same incident as press release "
        "english/press/20261012_12665.html (preserved bean curd, Bacillus cereus 140,000/g), "
        "which stays published. Resurrected by the Recalls row-merge; removed again. " + OP),
    "whatsnew_fa/2026_632.html": (
        "duplicate: CFS (HK) Food Alert 2026_632 is the same incident as press release "
        "english/press/20261009_12660.html (Marks & Spencer salad and pesto, Salmonella), "
        "which stays published. Resurrected by the Recalls row-merge; removed again. " + OP),
}
WITHDRAW = (f" [SUPERSEDED WITHDRAWN {TODAY} — the published copy was the error: the "
            "refusal stands; the row is removed from Recalls. {OP}]").replace("{OP}", OP)

FICHE = "https://rappel.conso.gouv.fr/fiche-rappel/23785/interne"
SHRIMP = {
    "Company": "CRUSTA C",
    "Product": ("Whole cooked chilled shrimp, 20/30 count, wild Penaeus indicus caught in the "
                "Western Indian Ocean (FAO 51), lot 9262800232, use by 16/10/2026"),
}
SHRIMP_NOTE = (f" {OP} RappelConso fiche 23785 (2026-10-0114), published 16/10/2026, read on "
               "its own page 2026-10-18: origine de la fiche \"CRUSTA C\"; lot 9262800232; DLC "
               "16/10/2026; sold 07/10–13/10/2026. [original product: Crevette entière cuite "
               "réfrigérée 20/30 sauvage penaeus indicus péchée en océan indien ouest FAO51] "
               "[collector product was 'cr20cz', a scraper artifact]")


def _head(ws):
    return [c.value for c in ws[1]]


def _key(url: str):
    u = (url or "").strip().lower()
    return next((k for k in REMOVE if k in u), None)


def remove_rows(wb) -> int:
    ws, arch = wb["Recalls"], wb["Rejected"]
    head, ah = _head(ws), _head(arch)
    ucol = head.index("URL") + 1
    n = 0
    for r in range(ws.max_row, 1, -1):
        k = _key(str(ws.cell(r, ucol).value or ""))
        if not k:
            continue
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(head)}
        out = []
        for h in ah:
            if h in ("RejectReason", "RejectionReason"):
                out.append(REMOVE[k])
            elif h == "RejectedBy":
                out.append("operator review " + TODAY)
            elif h == "RejectedAt":
                out.append(TODAY)
            elif h == "Status":
                out.append("rejected")
            else:
                out.append(row.get(h, ""))
        arch.append(out)
        ws.delete_rows(r, 1)
        n += 1
        print(f"  Recalls -> Rejected: {row.get('Date')} {row.get('Source')} {row.get('URL')}")
    for sheet in ("Weekly_Review",):
        if sheet in wb.sheetnames:
            wr = wb[sheet]
            wh = _head(wr)
            if "URL" in wh:
                wu = wh.index("URL") + 1
                for r in range(wr.max_row, 1, -1):
                    if _key(str(wr.cell(r, wu).value or "")):
                        wr.delete_rows(r, 1)
    return n


def withdraw_superseded(wb) -> int:
    n = 0
    for sheet, rc in (("Rejected", "RejectReason"), ("Weekly_Rejected", "RejectionReason")):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        col = {h: i for i, h in enumerate(_head(ws))}
        if "URL" not in col or rc not in col:
            continue
        for row in ws.iter_rows(min_row=2):
            if not _key(str(row[col["URL"]].value or "")):
                continue
            reason = str(row[col[rc]].value or "")
            if "superseded" in reason.lower() and "SUPERSEDED WITHDRAWN" not in reason:
                row[col[rc]].value = (reason + WITHDRAW).strip()
                n += 1
    return n


def split_rasff(wb) -> int:
    from pipeline._language import split_bilingual
    ws = wb["Recalls"]
    col = {h: i for i, h in enumerate(_head(ws))}
    n = 0
    for row in ws.iter_rows(min_row=2):
        if str(row[col["Source"]].value or "") != "RASFF (EU)":
            continue
        notes = []
        for c in ("Product", "Reason"):
            v = row[col[c]].value
            if not isinstance(v, str) or not v.strip():
                continue
            en = split_bilingual(v)
            if en and en != v:
                row[col[c]].value = en
                notes.append(f'[original {c.lower()} (bilingual RASFF subject): "{v}"]')
        if notes:
            row[col["Notes"]].value = (str(row[col["Notes"]].value or "") + f" {OP} "
                                       + " ".join(notes)).strip()
            n += 1
    return n


def fix_shrimp(wb) -> int:
    ws = wb["Pending"]
    col = {h: i for i, h in enumerate(_head(ws))}
    n = 0
    for row in ws.iter_rows(min_row=2):
        if str(row[col["URL"]].value or "").strip().lower() != FICHE:
            continue
        if row[col["Company"]].value == SHRIMP["Company"]:
            continue
        for k, v in SHRIMP.items():
            row[col[k]].value = v
        row[col["Notes"]].value = (str(row[col["Notes"]].value or "") + SHRIMP_NOTE).strip()
        n += 1
    return n


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    n_rm = remove_rows(wb)
    n_w = withdraw_superseded(wb)
    n_s = split_rasff(wb)
    n_f = fix_shrimp(wb)
    wb.save(XLSX)
    wb.close()
    print(f"SUPERSEDED stamps withdrawn: {n_w} · RASFF rows split: {n_s} · 23785 filled: {n_f}")
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print(f"ROWS_REMOVED={n_rm}")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
