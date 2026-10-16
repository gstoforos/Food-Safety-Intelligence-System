# -*- coding: utf-8 -*-
"""Operator decisions 2026-10-16. NOTHING IS REMOVED from Recalls.

Operator: "FSIS Grace Vienna sausages — no name; add RappelConso 23741."

A. USDA FSIS PHA-10062026-01 (2026-10-06), Grace chicken Vienna-style
   sausages. Held in Pending since 2026-10-13 because the row named no
   recalling firm. FSIS names none: the alert says only "the importer".
   Operator ruling: publish with no named firm. Company follows the
   register's existing wording for an unnamed firm ("Not named on the
   notice (...)", as on the 2026-10-07 and 2026-10-13 RappelConso rows).
   The alert page (fsis.usda.gov) returns a fetch error from this sandbox
   (Akamai); its wording was read 2026-10-16 as quoted by Food Safety News
   (https://www.foodsafetynews.com/2026/10/usda-issues-public-warning-about-vienna-sausage/):
   chicken Vienna-style sausages "illegally imported from Jamaica, a
   country that is not eligible to export meat or poultry products to the
   United States"; 140-gram cans of "GRACE READY TO EAT CHICKEN VIENNA
   SAUSAGES" and "... HOT & SPICY", any date, "EST. NO G-01" inside the
   Jamaica mark of inspection; distributed by the importer to retail in
   Alabama, Florida, Georgia, Maryland and Texas (matching the row's own
   FSIS API note); "no confirmed reports of illness or injury". The URL
   stays the FSIS notice. The Product held two bullet characters and the
   Reason was cut off mid-word ("public health aler"); both are rewritten.
   The Rejected copy of the row (marked UNRESOLVED 2026-10-12) is stamped
   SUPERSEDED so the re-promotion guard does not bar it.

B. RappelConso fiche 23741 (ref. 2026-10-0070), published 09/10/2026, read
   on its own page 2026-10-16: "RAISIN ROSE RED GLOBE CAT I ORIGINE
   ITALIE", brand "CONTE MAX", lot "52", sold "Du 22/09/2026 au
   24/09/2026", "France entière", distributors "SCANORMANDE-LECASUD-ALDI
   SAUVIAN", motif "dépassement LMR Abamectin", risk "Dépassement des
   limites autorisées de pesticides", symptoms "Nausées, douleurs,
   abdominales, maux de tête". The fiche prints no legal name for a
   manufacturer; its only firm is "Origine de la fiche : CARPENTIERE
   FRANCE", the submitter — recorded as such, not promoted to recalling
   firm. Verified absent from all sheets on 2026-10-14. It falls in the
   CLOSED week 41 (5–11 Oct), whose report and email are out; operator
   ruling: add it. The weekly-updates check rebuilds W41 and lists it in
   docs/data/weekly-updates-pending.json for the update mailer.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
OP = "[operator ruling 2026-10-16]"

FSIS_URL = ("https://www.fsis.usda.gov/recalls-alerts/fsis-issues-public-health-alert-"
            "ineligible-chicken-vienna-style-sausage-products")
FSN_URL = "https://www.foodsafetynews.com/2026/10/usda-issues-public-warning-about-vienna-sausage/"
GRACE = {
    "Company": "Not named on the notice (importer; brand Grace)",
    "Brand": "Grace",
    "Product": ("Grace Ready to Eat Chicken Vienna Sausages, 140 g cans — regular and "
                "Hot & Spicy, all dates (EST. NO G-01 inside the Jamaica mark of inspection)"),
    "Reason": ("Public health alert: chicken Vienna-style sausages illegally imported from "
               "Jamaica, a country not eligible to export meat or poultry products to the "
               "United States; distributed to retail in Alabama, Florida, Georgia, Maryland "
               "and Texas. No confirmed reports of illness. Do not eat; discard or return."),
}
GRACE_NOTE = (f" {OP} Published with no named firm: FSIS PHA-10062026-01 names only "
              f"\"the importer\". In scope as an import-eligibility case under the "
              f"operator ruling 2026-09-30 (\"in scope, as uninspected\"). SOURCES: "
              f"{FSIS_URL} (FSIS's own notice, the published URL; fsis.usda.gov answers "
              f"with a fetch error from the review sandbox); {FSN_URL} (alert wording "
              f"quoted, read 2026-10-16).")
GRACE_SUPERSEDE = (f" [SUPERSEDED 2026-10-16 — {OP} the alert names no firm; published "
                   "with Company 'Not named on the notice (importer; brand Grace)']")

FICHE = "https://rappel.conso.gouv.fr/fiche-rappel/23741/interne"
GRAPES = {
    "Date": "2026-10-09",
    "Source": "RappelConso (FR)",
    "Company": "Not named on the notice (fiche filed by Carpentiere France)",
    "Brand": "CONTE MAX",
    "Product": ("Red Globe red grapes, category I, origin Italy (Raisin rosé Red Globe cat. I "
                "origine Italie), lot 52, sold 22/09–24/09/2026"),
    "Pathogen": "Abamectin (pesticide residue)",
    "Reason": ("Abamectin above the maximum residue limit (MRL). Stated risk: authorized "
               "pesticide limits exceeded; possible nausea, abdominal pain, headache. Sold "
               "throughout France (Scanormande, Lecasud, Aldi Sauvian)."),
    "Class": "Recall",
    "Country": "France",
    "Outbreak": 0,
    "URL": FICHE,
    "Notes": (f"{OP} RappelConso fiche 23741 (2026-10-0070), published 09/10/2026, read on its "
              "own page 2026-10-16: motif \"dépassement LMR Abamectin\"; risque \"Dépassement "
              "des limites autorisées de pesticides\"; distributeurs \"SCANORMANDE-LECASUD-ALDI "
              "SAUVIAN\"; \"Origine de la fiche : CARPENTIERE FRANCE\" (submitter, no legal "
              "name printed). Added to the closed week 41 by operator ruling. "
              "[original product: RAISIN ROSE RED GLOBE CAT I ORIGINE ITALIE]"),
}


def _head(ws):
    return [c.value for c in ws[1]]


def fix_grace(wb) -> int:
    n = 0
    ws = wb["Pending"]
    col = {h: i for i, h in enumerate(_head(ws))}
    for row in ws.iter_rows(min_row=2):
        if str(row[col["URL"]].value or "").strip().lower() != FSIS_URL:
            continue
        if row[col["Company"]].value == GRACE["Company"]:
            continue
        for k, v in GRACE.items():
            row[col[k]].value = v
        row[col["Notes"]].value = (str(row[col["Notes"]].value or "") + GRACE_NOTE).strip()
        n += 1
    for sheet, rc in (("Rejected", "RejectReason"), ("Weekly_Rejected", "RejectionReason")):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        col = {h: i for i, h in enumerate(_head(ws))}
        if "URL" not in col or rc not in col:
            continue
        for row in ws.iter_rows(min_row=2):
            if str(row[col["URL"]].value or "").strip().lower() != FSIS_URL:
                continue
            reason = str(row[col[rc]].value or "")
            if "superseded" not in reason.lower():
                row[col[rc]].value = (reason + GRACE_SUPERSEDE).strip()
    return n


def add_grapes() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    present = False
    for sheet in ("Recalls", "Pending"):
        rows = wb[sheet].iter_rows(values_only=True)
        hdr = list(next(rows))
        i = hdr.index("URL")
        present |= any(r and str(r[i] or "").strip().lower() == FICHE for r in rows)
    wb.close()
    if present:
        print("  23741 already in Recalls or Pending — left alone")
        return 0
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as fh:
        json.dump([GRAPES], fh, ensure_ascii=False)
        path = fh.name
    out = subprocess.run([sys.executable, str(ROOT / "tools" / "add_manual_row.py"), path],
                         cwd=str(ROOT), capture_output=True, text=True)
    print((out.stdout or "").rstrip()[-800:])
    if out.returncode != 0:
        print("  add_manual_row FAILED:", (out.stderr or "")[-800:])
        return 0
    return 1


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    n = fix_grace(wb)
    wb.save(XLSX)
    wb.close()
    print(f"A — FSIS Grace PHA filled for publication: {n}")
    print(f"B — RappelConso 23741 queued: {add_grapes()}")
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print("ROWS_REMOVED=0")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
