# -*- coding: utf-8 -*-
"""Review of the 2026-10-17 morning fix. NOTHING IS REMOVED from Recalls
by this script (repair_2026_10_17.py removes the two NAFDAC rows).

Rebuilding the morning fix on main at 7e4a99c5, promote_gate_passing would
have published four rows that reached Pending after the morning pass. Each
passed every publish-gate rule; each is archived here to Weekly_Rejected
(append-only), and pipeline/_publish_gate.py rule 2d now refuses the first
two shapes:

  1. NAFDAC (NG) public-alert-no-041-2024 — Nivea roll-on DEODORANT, a 2024
     EU Safety Gate (RAPEX) alert on dangerous NON-FOOD products, Date set
     to the 2026 discovery day. Not food, and two years old.
  2. AGES (AT) English page .../en/human/product-warnings-product-recalls/
     produkt/produktwarnung-salmonellen-in-tomaten-aus-zwei-regionen-in-
     sizilien — Company "Product Warnings & Product" (the site heading),
     Product "AGES". The German page of the same warning (same slug) is the
     published row; this is its duplicate.
  3. CFS (HK) Food Alert whatsnew_fa/2026_633 — preserved bean curd, Bacillus
     cereus 140,000 per gram: the same incident as the CFS press release
     english/press/20261012_12665.html, already published (Kung Wo Dou Bun
     Chong). One incident, one row (as with the Camembert, 2026-10-11).
  4. CFS (HK) Food Alert whatsnew_fa/2026_632 — Marks & Spencer salad and
     pesto, Salmonella: the same incident as the CFS press release
     english/press/20261009_12660.html, already published.

  5. Every other Pending row whose URL carries an alert number from a
     different year than its Date (three NAFDAC rows arrived after the
     morning pass: 0047/2021 pediatric drug suspension, 034/2024 Philippine
     product warning, 027/2021 statins drug-safety notice). Old alerts dated
     on their discovery day; the publish gate (rule 2d) now refuses them,
     so they are archived rather than left to sit in Pending.

  6. SZPI (CZ) sildenafil supplements, published 06. 10. 2026 (read on the
     SZPI page 2026-10-17): Product was the bare word "Supplement". The notice
     names two products — ELIXIRMAN, food supplement, 2.4 g (8 capsules of
     300 mg), printed date 14-03-2027, seller ELIXIRMAN s.r.o., made by
     Šang-čchiou Ltd, China; and Capatrex, 10 capsules, use-by 01/2028,
     distributor C.P.X. Karlovy Vary, made by Paramount EU, LLC (UK). Filled
     from the notice. The row falls in the closed week 41 (Date 2026-10-06).

Also: the published CFS (HK) bean curd row reads "Shum Shui Po"; the CFS
press release says "Sham Shui Po" (read 2026-10-13). The 2026-10-14 pass
re-applied an older copy of repair_2026_10_13.py that still carried the
misspelling. Corrected here.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
TODAY = "2026-10-17"
OP = "[operator review 2026-10-17]"

ARCHIVE = {
    "public-alert-no-041-2024-alert-on-the-recall-of-nivea-roll-on-deodorant":
        ("not_food: NAFDAC public alert 041/2024 — Nivea roll-on deodorant, a 2024 EU "
         "Safety Gate (RAPEX) alert on non-food products; a cosmetic, not a food, and "
         "dated 2026 only because that is when it was found. " + OP),
    "/en/human/product-warnings-product-recalls/produkt/produktwarnung-salmonellen-in-tomaten-aus-zwei-regionen-in-sizilien":
        ("duplicate: English page of the AGES warning on Salmonella Strathcona in Sicilian "
         "tomatoes, published from its German page (same slug); Company held the site "
         "heading 'Product Warnings & Product'. " + OP),
    "cfs.gov.hk/english/whatsnew/whatsnew_fa/2026_633.html":
        ("duplicate: CFS (HK) Food Alert on preserved bean curd, Bacillus cereus 140,000/g "
         "— the same incident as press release english/press/20261012_12665.html, "
         "already published. One incident, one row. " + OP),
    "cfs.gov.hk/english/whatsnew/whatsnew_fa/2026_632.html":
        ("duplicate: CFS (HK) Food Alert on Marks & Spencer salad and pesto, Salmonella — "
         "the same incident as press release english/press/20261009_12660.html, already "
         "published. One incident, one row. " + OP),
}
BEANCURD = "https://www.cfs.gov.hk/english/press/20261012_12665.html"
SZPI = ("https://www.szpi.gov.cz/clanek/potravinarska-inspekce-zjistila-dva-doplnky-stravy-"
        "s-nepovolenym-obsahem-sildenafilu-latky-k-lecbe-erektilni-disfunkce.aspx")
SZPI_PRODUCT = ("Food supplements with undeclared sildenafil: ELIXIRMAN, 2.4 g (8 capsules of "
                "300 mg), printed date 14-03-2027 (made by Šang-čchiou Ltd, China); Capatrex, "
                "10 capsules, use-by 01/2028 (distributor C.P.X. Karlovy Vary; made by "
                "Paramount EU, LLC, UK)")


from pipeline._publish_gate import _ALERT_NO_YEAR  # noqa: E402


def _head(ws):
    return [c.value for c in ws[1]]


def archive(wb) -> int:
    pend, arch = wb["Pending"], wb["Weekly_Rejected"]
    ph, ah = _head(pend), _head(arch)
    ucol = ph.index("URL") + 1
    n = 0
    for r in range(pend.max_row, 1, -1):
        url = str(pend.cell(r, ucol).value or "").strip().lower()
        why = next((v for k, v in ARCHIVE.items() if k in url), None)
        if not why:
            m = _ALERT_NO_YEAR.search(url)
            d = str(pend.cell(r, ph.index("Date") + 1).value or "")[:4]
            if m and d.isdigit() and m.group(2) != d:
                why = (f"old alert: the URL's alert number {m.group(0)!r} is from "
                       f"{m.group(2)}, the Date ({d}) is the day it was found — not a "
                       f"{d} recall; NAFDAC's discovery channel re-surfaces old alerts. " + OP)
        if not why:
            continue
        row = {h: pend.cell(r, i + 1).value for i, h in enumerate(ph)}
        out = []
        for h in ah:
            if h == "RejectionReason":
                out.append(why)
            elif h == "RejectedBy":
                out.append("operator review " + TODAY)
            elif h == "Reviewed":
                out.append("Y")
            elif h == "Week_Added":
                out.append(TODAY)
            else:
                out.append(row.get(h, ""))
        arch.append(out)
        pend.delete_rows(r, 1)
        n += 1
        print(f"  Pending -> Weekly_Rejected: {url}")
    return n


def fix_sham(wb) -> int:
    ws = wb["Recalls"]
    h = _head(ws)
    u, rcol = h.index("URL"), h.index("Reason")
    n = 0
    for row in ws.iter_rows(min_row=2):
        if str(row[u].value or "").strip() == BEANCURD and "Shum Shui Po" in str(row[rcol].value or ""):
            row[rcol].value = str(row[rcol].value).replace("Shum Shui Po", "Sham Shui Po")
            n += 1
    return n


def fix_szpi(wb) -> int:
    n = 0
    for sheet in ("Pending", "Recalls"):
        ws = wb[sheet]
        h = _head(ws)
        u, pc, nc = h.index("URL"), h.index("Product"), h.index("Notes")
        for row in ws.iter_rows(min_row=2):
            if str(row[u].value or "").strip() == SZPI and str(row[pc].value or "") == "Supplement":
                row[pc].value = SZPI_PRODUCT
                row[nc].value = (str(row[nc].value or "") + f" {OP} Product filled from the SZPI "
                                 "notice (published 06. 10. 2026), read 2026-10-17; "
                                 "[original product: Supplement]").strip()
                n += 1
    return n


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    a, s, z = archive(wb), fix_sham(wb), fix_szpi(wb)
    print(f"SZPI product filled: {z}")
    if a or s or z:
        wb.save(XLSX)
    wb.close()
    print(f"archived from Pending: {a} · Sham Shui Po corrected: {s}")
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print("ROWS_REMOVED=0")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
