# -*- coding: utf-8 -*-
"""Operator review 2026-10-28.

A. FSAI (IE) 2026-06-12, Surasang Frozen Oysters, Norovirus — published at
   Tier 3. Operator ruling: every other norovirus row in the register is
   Tier 2 (Tier 1 only where the regulator links illnesses), so this one is
   Tier 2. Only Tier and a Notes stamp change.
B. Two CAA (JP) recalls from the daily-global brief of 2026-10-28, each read
   on its own recall.caa.go.jp detail page today. The detail pages carry no
   掲載日, so Date is the page's own 対応開始日 (response start date).
   Queued through tools/add_manual_row.py; normal review path.
NOTHING IS REMOVED FROM Recalls.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
TODAY = "2026-10-28"

SURASANG_URL = ("https://www.fsai.ie/news-and-alerts/food-alerts/"
                "recall-of-a-batch-of-surasang-frozen-oysters")
SURASANG_NOTE = ("[operator review 2026-10-28: Tier 3 -> 2. Operator ruling: "
                 "norovirus is Tier 2 as for every other norovirus row in the "
                 "register (Tier 1 only where the regulator links illnesses).]")

NEW_ROWS = [
    {
        "Date": "2026-10-26",
        "Source": "CAA (JP)",
        "Company": "Starbucks Coffee Japan",
        "Brand": "Starbucks",
        "Product": "Chocolate macaron",
        "Pathogen": "Foreign material (mold-like matter)",
        "Reason": "Possible contamination by a mold-like substance. "
                  "Refund / recall.",
        "Class": "Recall",
        "Country": "Japan",
        "Outbreak": 0,
        "URL": "https://www.recall.caa.go.jp/result/detail.php?"
               "rcl=00000036052&screenkbn=01",
        "Notes": (
            "[operator review 2026-10-28: daily-global MISSING item. Read on "
            "the CAA detail page: 件名 「スターバックスコーヒージャパン「チョコ"
            "レートマカロン」 - 返金／回収」, 回収理由 「カビ様のものによる汚"
            "染の可能性」, 対応開始日 2026年10月26日. No lot, date mark or "
            "region is printed. The page shows no 掲載日, so Date is the "
            "対応開始日. Pathogen filed as for the CAA Wajima jelly row "
            "('mold-like matter').] [original product: チョコレートマカロン]"),
    },
    {
        "Date": "2026-10-25",
        "Source": "CAA (JP)",
        "Company": "Takusu",
        "Brand": "Takusu",
        "Product": "Soy protein powder, sweet potato flavor, 1 kg "
                   "(sold online by ezobolic)",
        "Pathogen": "Foreign material (plastic fragments)",
        "Reason": "Foreign matter: fragments of the enclosed plastic scoop "
                  "found in the product. Refund / recall.",
        "Class": "Recall",
        "Country": "Japan",
        "Outbreak": 0,
        "URL": "https://www.recall.caa.go.jp/result/detail.php?"
               "rcl=00000036055&screenkbn=01",
        "Notes": (
            "[operator review 2026-10-28: daily-global MISSING item. Read on "
            "the CAA detail page: 件名 「タクス「ソイプロテインスイートポテト"
            "風味1kg」（ezobolicのオンラインで販売） - 返金／回収」, 回収理由 "
            "「異物（添付のプラスチック製のスプーンの破片）混入」, 対応開始日 "
            "2026年10月25日. Company is the business named first in the 件名; "
            "the contact field is blank. No lot, date mark or region is "
            "printed. The page shows no 掲載日, so Date is the 対応開始日.] "
            "[original product: ソイプロテインスイートポテト風味1kg]"),
    },
]


def _head(ws):
    return [c.value for c in ws[1]]


def _retier_surasang() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Recalls"]
    h = _head(ws)
    ui, ti, ni = (h.index("URL") + 1, h.index("Tier") + 1,
                  h.index("Notes") + 1)
    n = 0
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, ui).value or "").strip().rstrip("/") != SURASANG_URL:
            continue
        if str(ws.cell(r, ti).value) == "2":
            continue
        ws.cell(r, ti).value = 2
        notes = str(ws.cell(r, ni).value or "").strip()
        ws.cell(r, ni).value = (notes + " " + SURASANG_NOTE).strip()
        n += 1
        print("  Recalls: FSAI Surasang Frozen Oysters (Norovirus) Tier -> 2")
    wb.save(XLSX)
    wb.close()
    return n


def main() -> int:
    retiered = _retier_surasang()
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False,
                                     encoding="utf-8") as fh:
        json.dump(NEW_ROWS, fh, ensure_ascii=False, indent=2)
        rows_path = fh.name
    try:
        proc = subprocess.run(
            [sys.executable, str(ROOT / "tools" / "add_manual_row.py"),
             rows_path], cwd=str(ROOT), text=True, capture_output=True)
        print(proc.stdout.rstrip())
        if proc.stderr.strip():
            print(proc.stderr.rstrip(), file=sys.stderr)
    finally:
        Path(rows_path).unlink(missing_ok=True)
    if proc.returncode != 0:
        raise SystemExit(f"add_manual_row refused a row ({proc.returncode})")
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print(f"ROWS_RETIERED={retiered}")
    print(f"ROWS_QUEUED={len(NEW_ROWS)}")
    print("ROWS_REMOVED=0")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
