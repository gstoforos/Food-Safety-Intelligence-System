# -*- coding: utf-8 -*-
"""Operator review 2026-10-27: two invented EFET rows, before they publish.

Both came from the Greek news gap finder (discovered via skai.gr) and both
carry Company "Μπιφτέκι" — the Greek word for "burger", not a firm — and the
invented identifier "(Recall ID EFET-2026-12345)":

  * Date 2026-10-26, Product "Cured pork belly", Listeria, URL a mangled
    copy of EFET item 5379 ("…/item/5379-%20deltio-typou-acnaklisi-…").
    EFET item 5379 is the April 2026 recall of Nikolaos Tsatsoulis & Sons
    feta PDO for Listeria monocytogenes, ALREADY PUBLISHED in Recalls under
    its correct URL (Date 2026-04-09). Nothing about pork belly or
    26 October is on that notice. It passed every publish-gate rule and the
    promoter was about to publish it.
  * Date 2026-10-22, Company "Μπιφτέκι", Brand "ΜΑΓΓΙ", Product "Chicken
    mince", Salmonella, EFET item 5390 — the same row removed from Recalls
    on 2026-10-23 (repair_2026_10_23d) as a duplicate of MEATKA item 5424.
    The gap finder re-queued it; a previously rejected key is replaced for
    re-validation.

Both are archived to Rejected (append-only, reason recorded) from wherever
they sit — Pending, or Recalls if an hourly promotion reached them first
(hence remove-rows in the commit message). Weekly_Review copies of a row
that never publishes are left to the weekly capture.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
TODAY = "2026-10-27"

BAD_COMPANY = "Μπιφτέκι"
REASONS = {
    "/item/5379-": (
        "REJECTED — invented row: Company 'Μπιφτέκι' is the Greek word for "
        "'burger', Product 'Cured pork belly', Date 2026-10-26 and '(Recall ID "
        "EFET-2026-12345)' appear nowhere on EFET item 5379, which is the "
        "April 2026 Tsatsoulis feta PDO Listeria recall already published "
        "under its correct URL. News-discovered (skai.gr). Operator review "
        + TODAY + "."),
    "/item/5390-": (
        "REJECTED — duplicate re-queued: the EFET item 5390 row removed from "
        "Recalls on 2026-10-23 as a duplicate of MEATKA item 5424 (Company "
        "'Μπιφτέκι' is the Greek word for 'burger'; recall id invented). "
        "Operator review " + TODAY + "."),
}


def _head(ws):
    return [c.value for c in ws[1]]


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    rej = wb["Rejected"]
    rh = _head(rej)
    moved = removed = 0
    for sheet in ("Pending", "Recalls"):
        ws = wb[sheet]
        h = _head(ws)
        ui, ci = h.index("URL") + 1, h.index("Company") + 1
        for r in range(ws.max_row, 1, -1):
            url = str(ws.cell(r, ui).value or "")
            if "efet.gr" not in url:
                continue
            if str(ws.cell(r, ci).value or "").strip() != BAD_COMPANY:
                continue
            reason = next((v for k, v in REASONS.items() if k in url), None)
            if not reason:
                continue
            row = {k: ws.cell(r, i + 1).value for i, k in enumerate(h)}
            rej.append([reason if k in ("RejectReason", "RejectionReason")
                        else "operator review " + TODAY if k == "RejectedBy"
                        else TODAY if k == "RejectedAt"
                        else "rejected" if k == "Status"
                        else row.get(k, "") for k in rh])
            ws.delete_rows(r, 1)
            moved += 1
            removed += sheet == "Recalls"
            print(f"  {sheet} -> Rejected: {row.get('Date')} "
                  f"{row.get('Product')!r} {url[:90]}")
    wb.save(XLSX)
    wb.close()
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print(f"ROWS_ARCHIVED={moved}")
    print(f"ROWS_REMOVED={removed}")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
