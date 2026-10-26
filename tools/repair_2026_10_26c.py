# -*- coding: utf-8 -*-
"""Operator review 2026-10-26 (afternoon): foreign material is Tier 3.

Operator rule: "foreign material glass etc is Tier 3; not Tier 2 or Tier 1".
pipeline/_pathogen_scope.enforce_tier1 now enforces it (the writer runs that
guard on every Recalls row at every save). This applies the same guard, and
nothing else, to the rows already in Recalls / Pending / Weekly_Review whose
Pathogen is physical-only. On 2026-10-26: 24 published rows off Tier 3
(20 at Tier 2, 4 at Tier 1, e.g. Ministero della Salute tiramisu "steel
sphere", FSANZ Arla pudding glass). Only Tier and a Notes stamp change.
NOTHING IS REMOVED FROM Recalls.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"


def main() -> int:
    import openpyxl
    from pipeline._pathogen_scope import _is_physical_only, enforce_tier1
    wb = openpyxl.load_workbook(XLSX)
    n = 0
    for sheet in ("Recalls", "Pending", "Weekly_Review"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        head = [c.value for c in ws[1]]
        if not {"Pathogen", "Tier", "Notes"} <= set(head):
            continue
        ti, ni = head.index("Tier") + 1, head.index("Notes") + 1
        for r in range(2, ws.max_row + 1):
            row = {h: ws.cell(r, i + 1).value for i, h in enumerate(head)}
            if not _is_physical_only(row.get("Pathogen")):
                continue
            before = row.get("Tier")
            enforce_tier1(row)
            if str(row.get("Tier")) == str(before):
                continue
            ws.cell(r, ti).value = row["Tier"]
            ws.cell(r, ni).value = row.get("Notes")
            n += 1
            print(f"  {sheet}: Tier {before} -> {row['Tier']}: "
                  f"{row.get('Source')} {str(row.get('Date'))[:10]} "
                  f"{str(row.get('Product'))[:45]!r} ({row.get('Pathogen')})")
    wb.save(XLSX)
    wb.close()
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print(f"ROWS_RETIERED={n}")
    print("ROWS_REMOVED=0")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
