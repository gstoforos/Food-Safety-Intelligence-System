# -*- coding: utf-8 -*-
"""Operator review 2026-10-23 (George): the EFET chicken-burger recall of
22 Oct 2026 is published twice. remove-rows: ONE row leaves Recalls.

Kept: the operator-reviewed row on EFET's own notice item 5424 (MEATKA S.A.,
frozen chicken burger patties 850 g, lot 15626, Salmonella spp.).

Removed and archived to Rejected: the news-discovered row (kathimerini.gr)
published 2026-10-23 with
    Company  "Μπιφτέκι"   — Greek for "burger", not a firm
    Brand    "ΜΑΓΓΙ"      — Maggi, from a different EFET notice that day
                            (MAGGI noodles, a labelling recall, out of scope)
    Product  "Chicken mince"
    Reason   "Presence of Salmonella spp (Recall ID EFET-2026-12345)"
                          — no such recall ID exists on any EFET notice
    URL      EFET item 5390 — an older, different EFET press release.
Its fields describe no single notice, so it cannot be repaired in place.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
TODAY = "2026-10-23"
URL_KEY = "anakleiseis-cat/item/5390-"
REASON = ("REJECTED — duplicate with mixed-up fields: the EFET chicken-burger "
          "recall of 22 Oct 2026 is published from EFET's own notice item 5424 "
          "(MEATKA S.A.). This news-discovered copy carried Company 'Μπιφτέκι' "
          "(Greek for 'burger'), Brand 'ΜΑΓΓΙ' (from a different notice), an "
          "invented recall ID 'EFET-2026-12345' and the URL of an older, "
          "different EFET notice (item 5390). Operator review 2026-10-23.")


def _head(ws):
    return [c.value for c in ws[1]]


def _archive(arch, row, reason):
    out = []
    for h in _head(arch):
        if h in ("RejectReason", "RejectionReason"):
            out.append(reason)
        elif h == "RejectedBy":
            out.append("operator review " + TODAY)
        elif h == "RejectedAt":
            out.append(TODAY)
        elif h == "Status":
            out.append("rejected")
        else:
            out.append(row.get(h, ""))
    arch.append(out)


def _match(row) -> bool:
    return (URL_KEY in str(row.get("URL") or "").lower()
            and str(row.get("Source") or "").startswith("EFET")
            and str(row.get("Date") or "")[:10] == "2026-10-22")


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    ws, arch = wb["Recalls"], wb["Rejected"]
    head = _head(ws)
    n = 0
    for r in range(ws.max_row, 1, -1):
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(head)}
        if not _match(row):
            continue
        _archive(arch, row, REASON)
        ws.delete_rows(r, 1)
        n += 1
        print(f"  Recalls -> Rejected: {row.get('Company')!r} / {row.get('Product')!r}")
    for sheet in ("Weekly_Review", "Pending"):
        if sheet not in wb.sheetnames:
            continue
        wr = wb[sheet]
        h2 = _head(wr)
        if "URL" not in h2:
            continue
        for r in range(wr.max_row, 1, -1):
            row = {h: wr.cell(r, i + 1).value for i, h in enumerate(h2)}
            if _match(row):
                wr.delete_rows(r, 1)
                print(f"  {sheet}: duplicate copy deleted")
    if n != 1:
        raise SystemExit(f"expected exactly 1 Recalls row to remove, found {n}")
    wb.save(XLSX)
    wb.close()
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print(f"ROWS_REMOVED={n}")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
