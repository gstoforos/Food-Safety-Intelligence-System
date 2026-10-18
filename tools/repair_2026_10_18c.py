# -*- coding: utf-8 -*-
"""Week 42 finalization, 2026-10-18 (operator: "Finalize them"; Kayem link
sent by the operator = in scope). NOTHING IS REMOVED from Recalls.

A. USDA FSIS — Kayem Foods Inc. (Chelsea, MA) recalls approx. 1,584,361 lb
   of ready-to-eat meat and poultry products "due to possible temperature
   abuse", announced Saturday 17 Oct 2026, Class I. In no sheet of the
   register: the FSIS collector never delivered it. The FSIS notice page
   returns a fetch error from the review sandbox (Akamai); its wording was
   read 2026-10-18 as quoted by WGCU and Food Safety News (sources in
   Notes): products "may have experienced temperature abuse during
   storage"; "EST. M/P 7839", "EST. P 20935A", "EST. M 9520"; produced
   approx. March 31 – Oct. 7, 2026; shipped to retail, wholesale club and
   foodservice locations across the United States and Hong Kong; found
   after a retail customer complaint about warm product at delivery on
   1 Oct 2026, traced to a warehouse cooler with temperature deviations;
   "no confirmed reports of illness". Operator ruling 2026-10-18:
   temperature abuse of ready-to-eat product is IN SCOPE.

B. RappelConso fiche 23769 (2026-10-14), published in French: Product and
   Reason translated faithfully, originals kept in Notes.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
OP = "[week-42 finalization 2026-10-18]"

KAYEM_URL = ("https://www.fsis.usda.gov/recalls-alerts/"
             "kayem-foods-inc--recalls-meat-and-poultry-products-due-possible-temperature-abuse")
KAYEM = {
    "Date": "2026-10-17",
    "Source": "USDA FSIS",
    "Company": "Kayem Foods Inc.",
    "Brand": "",
    "Product": ("Ready-to-eat meat and poultry products (about 1,584,361 lb, more than 100 "
                "items of various weights), produced approx. 31 Mar – 7 Oct 2026, bearing "
                "EST. M/P 7839, EST. P 20935A or EST. M 9520"),
    "Pathogen": "Temperature abuse (cold-chain deviation)",
    "Reason": ("Class I recall: products may have experienced temperature abuse during "
               "storage in a warehouse cooler that had temperature deviations; found after a "
               "retail complaint about warm product at delivery on 1 Oct 2026. No confirmed "
               "illnesses. Shipped to retail, wholesale club and foodservice locations "
               "nationwide in the US and to Hong Kong."),
    "Class": "Recall",
    "Country": "United States",
    "Outbreak": 0,
    "URL": KAYEM_URL,
    "Notes": (f"{OP} Operator ruling 2026-10-18: temperature abuse of ready-to-eat product "
              "is in scope (operator sent the FSIS notice URL). SOURCES: " + KAYEM_URL +
              " (FSIS's own notice, the published URL; fsis.usda.gov answers with a fetch "
              "error from the review sandbox); https://www.wgcu.org/health/2026-10-17/"
              "more-than-1-5-million-pounds-of-meat-poultry-products-recalled-nationwide and "
              "https://www.foodsafetynews.com/2026/10/more-than-1-5-million-pounds-of-hot-"
              "dogs-bacon-sausage-turkey-breasts-and-more-meat-products-of-various-brands-"
              "recalled/ (FSIS wording quoted, read 2026-10-18). Brand left empty: the "
              "products carry various brands listed only in FSIS's label PDF."),
}

FICHE = "https://rappel.conso.gouv.fr/fiche-rappel/23769/interne"
SPINACH = {
    "Product": "Leaf spinach, IQF, 50 g portions, 2.5 kg bag (epinards branche portion 50g iqf sachet x 2.5 kg)",
    "Reason": "Non-inert foreign bodies (bones and feathers)",
}


def translate(wb) -> int:
    ws = wb["Recalls"]
    col = {c.value: i for i, c in enumerate(ws[1])}
    n = 0
    for row in ws.iter_rows(min_row=2):
        if str(row[col["URL"]].value or "").strip().lower() != FICHE:
            continue
        notes = []
        for k, v in SPINACH.items():
            old = str(row[col[k]].value or "")
            if old != v:
                row[col[k]].value = v
                notes.append(f'[{"original product" if k == "Product" else "original Reason (fr)"}: "{old}"]')
        if notes:
            nc = col["Notes"]
            txt = str(row[nc].value or "").replace("[needs cleanup: Product not in English]", "").strip()
            row[nc].value = (txt + f" {OP} " + " ".join(notes)).strip()
            n += 1
    return n


def add_kayem() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    present = False
    for sheet in ("Recalls", "Pending"):
        rows = wb[sheet].iter_rows(values_only=True)
        hdr = list(next(rows))
        i = hdr.index("URL")
        present |= any(r and str(r[i] or "").strip().lower() == KAYEM_URL for r in rows)
    wb.close()
    if present:
        print("  Kayem already in Recalls or Pending — left alone")
        return 0
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as fh:
        json.dump([KAYEM], fh, ensure_ascii=False)
        path = fh.name
    out = subprocess.run([sys.executable, str(ROOT / "tools" / "add_manual_row.py"), path],
                         cwd=str(ROOT), capture_output=True, text=True)
    print((out.stdout or "").rstrip()[-600:])
    if out.returncode != 0:
        print("  add_manual_row FAILED:", (out.stderr or "")[-600:])
        return 0
    return 1


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    n = translate(wb)
    wb.save(XLSX)
    wb.close()
    print(f"23769 translated: {n}")
    print(f"Kayem queued: {add_kayem()}")
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print("ROWS_REMOVED=0")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
