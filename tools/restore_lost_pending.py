#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Restore the six Pending rows lost to a stale-workbook upload.

WHAT HAPPENED (2026-09-23)
==========================
A data-fix workbook cut against commit 2e89ff29 was uploaded after a daily
update had already landed. The upload replaced the whole file, so the six
Pending rows the daily update had added went with it:

    Pending at 0a4028be (the daily update) ... 19 rows
    Pending after the upload ................ 13 rows

Recalls was NOT touched — it stayed at 1764 — so nothing PUBLISHED was
lost. The casualties are all candidates awaiting review: four RappelConso
(FR), one CFIA (CA), one FDA (US).

This is the same failure as commit 02d97dde, which cost 29 published rows,
and it has the same cause: a whole-file upload carries the state of the
moment it was CUT, not the moment it is uploaded.

WHY A SCRIPT AND NOT ANOTHER WORKBOOK
=====================================
Cutting a corrected recalls.xlsx would repeat the mistake. The register is
written by workflows several times an hour, so any workbook produced here
is stale before it can be uploaded. This script instead reads whatever
recalls.xlsx it finds AT RUN TIME and appends only what is missing.

SAFE TO RUN TWICE. It matches on (Company, Product, URL) against every
sheet, appends nothing that is already present anywhere in the register,
and deletes nothing. Run it, or do not — the six rows may also come back
on their own, since RappelConso, CFIA and FDA all re-scrape recent
listings daily and nothing blocks re-entry.

USAGE
=====
    python tools/restore_lost_pending.py            # report only
    python tools/restore_lost_pending.py --apply    # write
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"

#: Captured from commit 0a4028be, "FSIS daily update 2026-09-23
#: (+0 approved, 19 pending)" — the last state in which they existed.
LOST_ROWS = [{'Date': '2026-09-23',
  'Source': 'RappelConso (FR)',
  'Company': 'Carrefour le Marché',
  'Brand': 'Carrefour le Marché',
  'Product': 'jeunes pousses carrefour le marché - sachet 125g',
  'Pathogen': 'Salmonella',
  'Reason': 'Detection of Salmonella',
  'Class': 'Voluntary',
  'Country': 'France',
  'Region': 'Europe',
  'Tier': 1,
  'Outbreak': 0,
  'URL': 'https://rappel.conso.gouv.fr/fiche-rappel/23588/interne',
  'Notes': 'carrefour',
  'ScrapedAt': '2026-09-23T15:07:21Z',
  'Status': 'pending',
  'RejectedBy': None},
 {'Date': '2026-09-23',
  'Source': 'RappelConso (FR)',
  'Company': 'U',
  'Brand': 'U',
  'Product': 'flan rhum raisins u différents conditionnements (1 ou '
             'plusieurs parts)',
  'Pathogen': None,
  'Reason': 'Suspicion de présence de corps étrangers métalliques',
  'Class': 'Voluntary',
  'Country': 'France',
  'Region': 'Europe',
  'Tier': 3,
  'Outbreak': 0,
  'URL': 'https://rappel.conso.gouv.fr/fiche-rappel/23587/interne',
  'Notes': 'magasins u',
  'ScrapedAt': '2026-09-23T15:07:21Z',
  'Status': 'pending_enrichment',
  'RejectedBy': None},
 {'Date': '2026-09-23',
  'Source': 'RappelConso (FR)',
  'Company': 'Pépite',
  'Brand': 'Pépite',
  'Product': 'pot 300g',
  'Pathogen': 'Salmonella',
  'Reason': 'Suspicion de Salmonella',
  'Class': 'Voluntary',
  'Country': 'France',
  'Region': 'Europe',
  'Tier': 1,
  'Outbreak': 0,
  'URL': 'https://rappel.conso.gouv.fr/fiche-rappel/23590/interne',
  'Notes': 'magasins bio',
  'ScrapedAt': '2026-09-23T15:07:21Z',
  'Status': 'pending',
  'RejectedBy': None},
 {'Date': '2026-09-23',
  'Source': 'RappelConso (FR)',
  'Company': 'Giffaud / les Delices de Clobert',
  'Brand': 'Giffaud / les Delices de Clobert',
  'Product': 'paupiette de poulet à la montagnarde',
  'Pathogen': 'Salmonella',
  'Reason': 'Presence of Salmonella',
  'Class': 'Voluntary',
  'Country': 'France',
  'Region': 'Europe',
  'Tier': 1,
  'Outbreak': 0,
  'URL': 'https://rappel.conso.gouv.fr/fiche-rappel/23600/interne',
  'Notes': 'leclerc',
  'ScrapedAt': '2026-09-23T15:07:21Z',
  'Status': 'pending',
  'RejectedBy': None},
 {'Date': '2026-09-22',
  'Source': 'CFIA',
  'Company': None,
  'Brand': 'R.J. King Fisheries Ltd. and Moncton Fish Market',
  'Product': 'Frozen cooked lobster meat',
  'Pathogen': 'Staphylococcus enterotoxin',
  'Reason': 'R.J. King Fisheries Ltd. and Moncton Fish Market brand frozen '
            'cooked lobster meat recalled due to Staphylococcus aureus',
  'Class': 'Class 2',
  'Country': 'Canada',
  'Region': 'North America',
  'Tier': 2,
  'Outbreak': 0,
  'URL': 'https://recalls-rappels.canada.ca/en/alert-recall/rj-king-fisheries-ltd-and-moncton-fish-market-brand-frozen-cooked-lobster-meat',
  'Notes': 'CFIA open-data NID=82632',
  'ScrapedAt': '2026-09-23T15:07:21Z',
  'Status': 'pending',
  'RejectedBy': None},
 {'Date': '2026-08-23',
  'Source': 'FDA',
  'Company': 'International Sprout Holdings, Inc',
  'Brand': '—',
  'Product': 'Alfalfa, Net Weight 50 lbs. bag, Product of USA, '
             'International Specialty Supply 1011 Volunteer Dr. '
             'Cookeville, TN 38506',
  'Pathogen': 'Salmonella',
  'Reason': 'Potential E. coli and Salmonella contamination',
  'Class': 'Class I',
  'Country': 'United States',
  'Region': 'North America',
  'Tier': 1,
  'Outbreak': 0,
  'URL': 'https://www.fda.gov/safety/recalls-market-withdrawals-safety-alerts?search_api_fulltext=H-1341-2026',
  'Notes': 'openFDA recall#=H-1341-2026; ev_id=99822; distrib=Wholesale '
           'consignees in 16 states: CA, FL, HI, IL, MA, MN, MO, MT, NY, '
           'OH, OR, PA, TN, SC, TX, UT; Puerto Rico; Canada,',
  'ScrapedAt': '2026-09-23T15:07:21Z',
  'Status': 'pending',
  'RejectedBy': None}]


def _key(company, product, url):
    return (str(company or "").strip().lower(),
            str(product or "").strip().lower()[:60],
            str(url or "").strip().lower())


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true",
                    help="write the rows (default: report only)")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("openpyxl is required", file=sys.stderr)
        return 2

    if not XLSX.exists():
        print(f"no register at {XLSX}", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(XLSX)

    # Present ANYWHERE counts as present: a row that has since been
    # promoted to Recalls or rejected must not be re-queued for review.
    present = {}
    for name in wb.sheetnames:
        ws = wb[name]
        head = [str(c.value or "") for c in ws[1]]
        cols = {c: head.index(c) for c in ("Company", "Product", "URL")
                if c in head}
        if len(cols) < 3:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            present[_key(row[cols["Company"]], row[cols["Product"]],
                         row[cols["URL"]])] = name

    ws = wb["Pending"]
    head = [str(c.value or "") for c in ws[1]]

    todo, skip = [], []
    for r in LOST_ROWS:
        k = _key(r.get("Company"), r.get("Product"), r.get("URL"))
        (skip if k in present else todo).append((r, present.get(k)))

    for r, where in skip:
        print(f"  already in {where:16} | {str(r.get('Company'))[:30]:30} "
              f"| {str(r.get('Product'))[:40]}")
    for r, _ in todo:
        print(f"  WOULD RESTORE            | {str(r.get('Company'))[:30]:30} "
              f"| {str(r.get('Product'))[:40]}")

    print(f"\n  {len(todo)} to restore, {len(skip)} already present")
    if not todo:
        return 0
    if not args.apply:
        print("  (report only — re-run with --apply to write)")
        return 0

    for r, _ in todo:
        # Build by HEADER NAME, never by position: the Pending sheet's
        # column order is not guaranteed to match the order these rows
        # were captured in, and writing by position is how a URL ends up
        # in the Notes column.
        ws.append([r.get(col) for col in head])
    wb.save(XLSX)
    print(f"  wrote {len(todo)} rows to {XLSX}")
    print("  UPLOAD THIS IMMEDIATELY — it is stale the moment a workflow runs")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
