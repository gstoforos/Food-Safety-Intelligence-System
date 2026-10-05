# -*- coding: utf-8 -*-
"""Accuracy-brief follow-up, 2026-10-05 (13:09 Athens). Four rows the
daily accuracy brief flagged; every value below was read on the authority
page named beside it. Idempotent: a row already fixed is left alone.

  1. SZPI (CZ) MASO WEST — Product "Italian salad" is a literal rendering of
     "Vlašský salát", a Czech potato salad; it misdescribes the food.
     szpi.gov.cz/clanek/varovani-pro-spotrebitele-listerie-ve-vlasskem-salatu.aspx
     (read 2026-10-05): 140 g / 300 g / 1 kg, batch L862026265, best before
     06.10.2026, MASO WEST s.r.o., Klatovy.
  2. RappelConso 23696 SODIMAZ LECLERC — Brand "—" although the row's own
     Notes say the fiche reads "Sans marque" and Brand was left empty.
  3. FSAI Wrights of Marino — Reason was a headline fragment ("due to the
     presence of ..."); Product lacked the pack and batch. fsai.ie alert
     2026.61 (read 2026-10-04): 400 g, batch Yellow A38, use-by 06/10/2026,
     Category 1 - For Action.
  4. NEWS — Food Safety News "Czech testing finds Salmonella problem in
     poultry meat" tagged Event=Outbreak on a retail-testing story whose
     title names no human case (rule fixed in scrapers/news.py).

Run from the repo root:  python tools/repair_2026_10_05b.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
import pipeline.merge_master as m  # noqa: E402

XLSX = Path("docs/data/recalls.xlsx")
TODAY = "2026-10-05"
SZPI = "https://www.szpi.gov.cz/clanek/varovani-pro-spotrebitele-listerie-ve-vlasskem-salatu.aspx"
SODIMAZ = "https://rappel.conso.gouv.fr/fiche-rappel/23696/interne"
WRIGHTS = "https://www.fsai.ie/news-and-alerts/food-alerts/recall-of-a-batch-of-wrights-of-mario-oak-sliced-s"
NEWS_LINK = "https://www.foodsafetynews.com/2026/10/czech-testing-finds-salmonella-problem-in-poultry-meat"


def _stamp(r, text):
    r["Notes"] = (str(r.get("Notes") or "").strip() + f" [accuracy-brief fix {TODAY}: {text}]").strip()
    r["LastUpdated"] = TODAY


def main() -> int:
    rec = m._load_sheet(XLSX, "Recalls", m.RECALLS_SCHEMA)
    pen = m._load_sheet(XLSX, "Pending", m.PENDING_SCHEMA)
    by = {str(r.get("URL", "")).strip().lower(): r for r in rec}
    n = 0

    r = by.get(SZPI.lower())
    if r and r.get("Product") == "Italian salad":
        new = ("Vlašský salát (Czech potato salad), 140 g / 300 g / 1 kg, "
               "batch L862026265, best before 06.10.2026")
        _stamp(r, "[original product: Italian salad] — a literal rendering of 'Vlašský salát' "
                  "that misdescribes the food; pack, batch and best-before read on the SZPI notice")
        r["Product"] = new
        n += 1
        print("1 MASO WEST product ->", new)

    r = by.get(SODIMAZ.lower())
    if r and str(r.get("Brand") or "").strip() in ("—", "-", "–"):
        r["Brand"] = ""
        _stamp(r, "Brand '—' cleared; the fiche reads 'Sans marque'")
        n += 1
        print("2 SODIMAZ brand cleared")

    r = by.get(WRIGHTS.lower())
    if r and str(r.get("Reason") or "").startswith("due to"):
        r["Reason"] = ("Batch recalled due to the presence of Listeria monocytogenes "
                       "(FSAI alert notification 2026.61, Category 1 - For Action)")
        if r.get("Product") == "Sliced Smoked Salmon":
            r["Product"] = "Sliced Smoked Salmon, 400 g, batch Yellow A38, use-by 06/10/2026"
        _stamp(r, "Reason was a headline fragment; Reason and pack/batch read on FSAI alert "
                  "2026.61. The earlier 'claude-check please enrich' note is resolved")
        n += 1
        print("3 Wrights of Marino reason/product filled")

    if n:
        m.save_xlsx_with_pending(m.sort_rows(rec), m.sort_rows(pen), XLSX)

    wb = openpyxl.load_workbook(XLSX)
    if "NEWS" in wb.sheetnames:
        ws = wb["NEWS"]
        h = [c.value for c in ws[1]]
        if "Link" in h and "Event" in h:
            il, ie = h.index("Link"), h.index("Event")
            for row in ws.iter_rows(min_row=2):
                if str(row[il].value or "").strip() == NEWS_LINK and row[ie].value == "Outbreak":
                    row[ie].value = "News"
                    n += 1
                    print("4 NEWS Czech testing: Outbreak -> News")
            wb.save(XLSX)

    if n:
        m.mirror_json_from_xlsx(XLSX, Path("docs/data/recalls.json"))
    print(f"ROWS_REMOVED=0  changes={n}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
