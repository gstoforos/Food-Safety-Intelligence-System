# -*- coding: utf-8 -*-
"""Morning fix 2026-10-19. ONE ROW IS REMOVED from Recalls — the commit
message carries "remove-rows".

A. ONE ROW OUT OF RECALLS — a published non-recall.

   At 03:27 UTC today the offline promoter published, Tier-1, into week 42:

       Date      2026-10-15
       Source    Livsmedelsverket (SE)
       Company   '121'
       Brand     '121'
       Product   'Edsviken - Livsmedelsverket'
       Pathogen  'Escherichia coli (generic)'
       Reason    'E. coli'
       URL       .../sok---muselkontroll-oppna-och-stangda-
                 produktionsomraden/edsviken
       Notes     '[via DDGS gap-finder, deterministic extract]
                  [promote 2026-10-19: ... passes every publish-gate rule]'

   The URL was read on 2026-10-19: it is Sweden's shellfish production-AREA
   classification page for the Edsviken area, open with a restriction on
   oyster harvesting because of high E. coli in the water. It names no firm,
   no brand, no product and no lot, and announces no recall. A water-quality
   classification is not a regulatory action (R2) and a status page is not a
   per-recall notice (R3). Company and Brand held the number 121 and Product
   held the page's <title>. The same row is removed from Weekly_Review (the
   Sunday review slice), and archived to Rejected — nothing is deleted.

   The gate half of this is pipeline/_publish_gate.py rule 2e in the same
   upload, with tests/test_a_number_is_not_a_firm_and_the_authority_is_not_
   the_product.py. Rule 2b closed the ACRONYM half of this gap on 2026-10-11
   and recorded the reason it existed; the full-name half stayed open.

B. THREE PUBLISHED ROWS WITH THE AUTHORITY OR ITS PORTAL IN A NAME FIELD.
   Every replacement value is read off the row's OWN Notes; nothing is
   fetched and nothing is invented.

   * row 1026, BVL (DE), 2026-07-07 — Company = Brand =
     'lebensmittelwarnung.de', the portal that published the notice, not the
     dairy. The row's own Notes carry the 2026-07-14 reviewer's reading of
     the page: "Company (Hof Dannwisch) and Product (Salatkäse Kräuter)
     confirmed". Company := 'Hof Dannwisch'. Brand := '—': no brand is named
     anywhere on the row, and a value that is not there is left empty.

   * row 2019, FSANZ (AU), 2026-03-05 — Product carried FSANZ's own site-name
     suffix: "Fresh 'N' Frozen Favour-ita Grated Coconut 400g | Food
     Standards Australia New Zealand". recall_review_agent flagged this shape
     on 2026-09-02 and its note says plainly "Those two need their Product
     field cleaned"; it was never done. The row's own Notes hold the FSANZ
     <h1>: "Bajaria Global Pty Ltd T/a The Spice Merchants - Fresh 'N' Frozen
     Favour-ita Grated Coconut 400g". Only the suffix is removed.

   * row 2065, FSAI (IE), 2026-01-23 — the OTHER row from that same
     2026-09-02 note. Product was nothing but the suffix,
     '| Food Safety Authority of Ireland', and Company = Brand held the
     headline 'Danone infant formula and follow-on formula'. The row's own
     Notes carry the FSAI lede verbatim: "Danone is recalling specific
     batches of its infant formula and follow-on formula due to the potential
     presence of cereulide", and the 2026-05-07 reviewer's "Company (Danone)
     and Product (infant formula and follow...)".

C. FOUR SILENT CONTRADICTIONS IN THE ARCHIVE.
   tests/test_a_recall_is_not_both_published_and_rejected.py has been RED on
   main (it fails on an unmodified clone). Four Rejected rows share a URL
   with a published recall and their RejectReason is empty, so the archive
   says nothing about which copy wins. The verdict already exists: the
   2026-10-16 pass wrote a full SUPERSEDED explanation into each row's
   NOTES and left the column the test and the readers key on blank. The
   note is copied into RejectReason; no judgement is added.

D. FOUR NAFDAC ROWS OUT OF PENDING.
   tests/test_alert_number_year_matches_date.py has also been RED on main.
   Four Pending rows carry an alert number from another year in their own
   URL — 041/2024, 034/2024, 0047/2021, 027/2021 — against a 2026 Date, and
   every one of them is a medicine or a cosmetic, not food (R1):
   Nivea roll-on deodorant, unregistered drug products in the Philippines,
   Lefin Pediatric Suspension, and statins. The publish gate already refuses
   all four, so none could ever be promoted; they are moved to Rejected so
   the queue and the test stop carrying them.

   NOT DONE, and it is the operator's call: three more NAFDAC medicine rows
   (030/2026 levamisole, 038/2026 adulterated medicines, 023/2026
   counterfeit proguanil) sit in Pending blocked only on 'Company is empty'
   and 'Pathogen is empty' — repairable-looking blockers. A Product-only
   medicine marker at the gate would close that, but the obvious markers
   cannot go in the Product/Reason/URL blob the existing non-food rule uses:
   measured on the live register, 'medicinal product', 'medicines',
   'antibiotic' and 'pharmaceutical' hit four GOOD published food rows
   through their Reason — veterinary-medicine residues in chicken strips
   (RappelConso 23745) and in milk, antibiotic residue in scampi, and an
   Aptamil row. Those are legitimate chemical-contaminant recalls.

E. ONE PENDING ROW ENRICHED FROM THE AUTHORITY PAGE, so the offline promoter
   can publish it. Pending has held Krasnyi Oktyabr Inc.'s SECOND October
   Listeria recall since 2026-10-17 with Date and Pathogen empty and Product
   set to the FDA headline, stamped "FDA HTML fallback — claude-check needs
   to enrich Date+Pathogen". claude_check does not run. Read on FDA's own
   page 2026-10-19 (fda.gov is pre-authorized and the URL was returned by
   search):

       Company Announcement Date  October 9, 2026
       FDA Publish Date           October 16, 2026
       firm      Krasnyi Oktyabr Inc., Brooklyn, New York
       brand     Starfish Seafood
       product   Starfish Seafood Premium Cold Smoked Norwegian Salmon
                 Sliced, 227 g (8 oz), vacuum-sealed, frozen
       lot       D080-148, best before 27.11.2027
       UPC       4751048940099 · origin Latvia
       hazard    potential Listeria monocytogenes contamination
       states    New York, Florida, Georgia, New Jersey, Pennsylvania, Ohio
       illness   "No illnesses have been reported to date"

   This is a different notice from the already-published sibling (Recalls
   row 94, the Starfish Seafood TRIO, lot D139149161-153, 2026-10-09): a
   different product, lot, UPC and distribution list, on its own FDA URL.
   Date := 2026-10-16, the regulator's own publish date; the company
   announcement date goes to Notes. Nothing on the row is inferred.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
TODAY = "2026-10-19"
OP = "[morning fix 2026-10-19]"

# --- A ----------------------------------------------------------------------
EDSVIKEN = ("livsmedelsverket.se/foretagande-regler-kontroll/"
            "sa-kontrolleras-ditt-foretag/"
            "sok---muselkontroll-oppna-och-stangda-produktionsomraden/edsviken")
EDSVIKEN_REASON = (
    "not a recall: the URL is Livsmedelsverket's shellfish production-AREA "
    "classification page for Edsviken — the area is open with a restriction on "
    "oyster harvesting because of high E. coli in the water. Read 2026-10-19: it "
    "names no firm, no brand, no product and no lot and announces no regulatory "
    "action against any product. A water-quality classification is not a recall "
    "(R2) and a status page is not a per-recall notice (R3). Company and Brand "
    "held '121' and Product held the page <title>. " + OP)

# --- B ----------------------------------------------------------------------
NAME_FIXES = {
    "lebensmittelwarnung.de/___lebensmittelwarnung.de/meldungen/2026/07_juli/"
    "260707_11_sh_salatkaese": {
        "Company": "Hof Dannwisch",
        "Brand": "—",
        "_note": ("Company/Brand were 'lebensmittelwarnung.de', the portal that "
                  "published the notice, not the firm; Source already records the "
                  "portal. The firm is named on this row's own Notes by the "
                  "2026-07-14 reviewer that read the page: 'Company (Hof "
                  "Dannwisch)'. No brand is named anywhere on the row, so Brand "
                  "is left empty rather than filled with the firm. " + OP),
    },
    "foodstandards.gov.au/food-recalls/recall-alert/"
    "bajaria-global-pty-ltd-ta-spice-merchants-fresh-n-frozen-favour-ita": {
        "Product": "Fresh 'N' Frozen Favour-ita Grated Coconut 400g",
        "_note": ("Product carried FSANZ's own site-name suffix "
                  "' | Food Standards Australia New Zealand'. Only the suffix is "
                  "removed; the product name is the one on this row's own Notes, "
                  "from the FSANZ <h1>. Flagged as needing exactly this cleanup by "
                  "recall_review_agent on 2026-09-02. " + OP),
    },
    "fsai.ie/news-and-alerts/latest-news/"
    "danone-infant-formula-and-follow-on-formula-recall": {
        "Company": "Danone",
        "Brand": "Danone",
        "Product": "Infant formula and follow-on formula (specific batches)",
        "_note": ("Product was nothing but the site-name suffix "
                  "'| Food Safety Authority of Ireland' and Company/Brand held the "
                  "headline. Both values are on this row's own Notes, from the FSAI "
                  "lede: 'Danone is recalling specific batches of its infant "
                  "formula and follow-on formula due to the potential presence of "
                  "cereulide'. The second of the two rows recall_review_agent "
                  "flagged on 2026-09-02. " + OP),
    },
}

# --- C ----------------------------------------------------------------------
SUPERSEDED_URLS = (
    "fsis.usda.gov/recalls-alerts/shanghai-ravioli-corporation-recalls-"
    "not-ready-eat-frozen-buffalo-chicken-products",
    "fsis.usda.gov/recalls-alerts/sempio-food-services-inc--recalls-"
    "ready-eat-chicken-stew-products-imported-without",
    "fsis.usda.gov/recalls-alerts/el-eden-import-distributor-corp-recalls-"
    "ineligible-pork-cracklings-products-imported",
    "fsis.usda.gov/recalls-alerts/asian-america-trading-inc--recalls-"
    "ineligible-siluriformes-fish-products-imported",
)

# --- D ----------------------------------------------------------------------
NAFDAC_OUT = {
    "public-alert-no-041-2024": (
        "not_food: NAFDAC public alert 041/2024 — Nivea roll-on DEODORANT, a "
        "cosmetic relayed from the EU Safety Gate for dangerous non-food products. "
        "Not food (R1), and the alert number in its own URL is from 2024 against a "
        "2026 Date. " + OP),
    "public-alert-no-034-2024": (
        "not_food: NAFDAC public alert 034/2024 — unregistered DRUG products in the "
        "Philippines. Not food (R1); Product and Pathogen empty and the Reason "
        "('Undeclared almond allergen') belongs to a different notice entirely. "
        "Alert number from 2024 against a 2026 Date. " + OP),
    "public-alert-no-0047-2021": (
        "not_food: NAFDAC public alert 0047/2021 — Lefin Pediatric Suspension, a "
        "medicine. Not food (R1), and the alert number in its own URL is from 2021 "
        "against a 2026 Date. " + OP),
    "public-alert-no-027-2021": (
        "not_food: NAFDAC public alert 027/2021 — statins, a drug-safety "
        "communication relayed from the US FDA. Not food (R1), and the alert number "
        "in its own URL is from 2021 against a 2026 Date. " + OP),
}

# --- E ----------------------------------------------------------------------
SALMON_URL = ("fda.gov/safety/recalls-market-withdrawals-safety-alerts/"
              "krasnyi-oktyabr-inc-recalls-starfish-seafood-premium-cold-"
              "smoked-norwegian-salmon-sliced-227g8oz")
SALMON = {
    "Date": "2026-10-16",
    "Company": "Krasnyi Oktyabr Inc.",
    "Brand": "Starfish Seafood",
    "Product": ("Starfish Seafood Premium Cold Smoked Norwegian Salmon Sliced, "
                "227 g (8 oz) vacuum-sealed frozen packages, lot D080-148, best "
                "before 27.11.2027, UPC 4751048940099, origin Latvia"),
    "Pathogen": "Listeria monocytogenes",
    "Reason": ("Potential contamination with Listeria monocytogenes. Distributed to "
               "New York, Florida, Georgia, New Jersey, Pennsylvania and Ohio. No "
               "illnesses reported to date."),
    "Class": "Recall",
    "Country": "United States",
    "Region": "North America",
    "Tier": 1,
    "Outbreak": 0,
    "Status": "pending",
}
SALMON_NOTE = (
    " " + OP + " Read on FDA's own notice page 2026-10-19: FDA Publish Date "
    "October 16, 2026 (used as Date); Company Announcement Date October 9, 2026; "
    "firm Krasnyi Oktyabr Inc. of Brooklyn, New York; brand Starfish Seafood; "
    "227 g (8 oz) vacuum-sealed frozen packages; lot D080-148; best before "
    "27.11.2027; UPC 4751048940099; origin Latvia; hazard potential Listeria "
    "monocytogenes; distributed to NY, FL, GA, NJ, PA, OH; 'No illnesses have been "
    "reported to date'. A DIFFERENT notice from the published sibling (Starfish "
    "Seafood TRIO, lot D139149161-153, 2026-10-09): different product, lot, UPC "
    "and distribution list, on its own FDA URL. The collector's Product was the "
    "page headline and Date/Pathogen were empty, stamped 'FDA HTML fallback — "
    "claude-check needs to enrich Date+Pathogen'; claude_check does not run.")


def _head(ws):
    return [c.value for c in ws[1]]


def _cols(ws):
    return {h: i for i, h in enumerate(_head(ws))}


def _archive(arch, row, head, reason):
    ah = _head(arch)
    out = []
    for h in ah:
        if h in ("RejectReason", "RejectionReason"):
            out.append(reason)
        elif h == "RejectedBy":
            out.append("operator review " + TODAY)
        elif h == "RejectedAt":
            out.append(TODAY)
        elif h == "Status":
            out.append("rejected")
        else:
            out.append(row.get(h, ""))
    arch.append(out)


def remove_edsviken(wb) -> int:
    """A — one non-recall out of Recalls, archived, and out of the review slice."""
    ws, arch = wb["Recalls"], wb["Rejected"]
    head = _head(ws)
    ucol = head.index("URL") + 1
    n = 0
    for r in range(ws.max_row, 1, -1):
        if EDSVIKEN not in str(ws.cell(r, ucol).value or "").lower():
            continue
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(head)}
        _archive(arch, row, head, EDSVIKEN_REASON)
        ws.delete_rows(r, 1)
        n += 1
        print(f"  Recalls -> Rejected: {row.get('Date')} {row.get('Source')} "
              f"{row.get('Product')!r}")
    for sheet in ("Weekly_Review", "Pending"):
        if sheet not in wb.sheetnames:
            continue
        wr = wb[sheet]
        col = _cols(wr)
        if "URL" not in col:
            continue
        for r in range(wr.max_row, 1, -1):
            if EDSVIKEN in str(wr.cell(r, col["URL"] + 1).value or "").lower():
                wr.delete_rows(r, 1)
                print(f"  {sheet}: Edsviken row deleted")
    return n


def fix_name_fields(wb) -> int:
    """B — the authority or its portal sitting in a name field."""
    ws = wb["Recalls"]
    col = _cols(ws)
    n = 0
    for row in ws.iter_rows(min_row=2):
        url = str(row[col["URL"]].value or "").lower()
        key = next((k for k in NAME_FIXES if k in url), None)
        if key is None:
            continue
        fix = NAME_FIXES[key]
        changed = False
        for field, value in fix.items():
            if field.startswith("_"):
                continue
            if row[col[field]].value != value:
                row[col[field]].value = value
                changed = True
        if not changed:
            continue
        row[col["Notes"]].value = (str(row[col["Notes"]].value or "") + " "
                                   + fix["_note"]).strip()
        row[col["LastUpdated"]].value = TODAY
        n += 1
        print(f"  Recalls name fix: {row[col['Source']].value} "
              f"{row[col['Date']].value}")
    return n


def annotate_superseded(wb) -> int:
    """C — copy the existing SUPERSEDED verdict into the empty reason column."""
    ws = wb["Rejected"]
    col = _cols(ws)
    n = 0
    for row in ws.iter_rows(min_row=2):
        url = str(row[col["URL"]].value or "").lower()
        if not any(u in url for u in SUPERSEDED_URLS):
            continue
        reason = str(row[col["RejectReason"]].value or "").strip()
        if "superseded" in reason.lower():
            continue
        notes = str(row[col["Notes"]].value or "")
        if "[SUPERSEDED" not in notes:
            continue
        verdict = notes[notes.index("[SUPERSEDED"):].strip()
        row[col["RejectReason"]].value = (
            "SUPERSEDED — the row published under this URL is the correct one; "
            "this archive entry is the audit trail of the vocabulary defect that "
            "refused it. The verdict was written into Notes by the 2026-10-16 pass "
            "and left out of this column; copied here verbatim, nothing added: "
            + verdict)
        n += 1
        print(f"  Rejected SUPERSEDED reason set: {row[col['Date']].value} "
              f"{row[col['Source']].value}")
    return n


def reject_nafdac(wb) -> int:
    """D — four non-food NAFDAC rows out of the promotion queue."""
    ws, arch = wb["Pending"], wb["Rejected"]
    head = _head(ws)
    ucol = head.index("URL") + 1
    n = 0
    for r in range(ws.max_row, 1, -1):
        url = str(ws.cell(r, ucol).value or "").lower()
        key = next((k for k in NAFDAC_OUT if k in url), None)
        if key is None:
            continue
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(head)}
        _archive(arch, row, head, NAFDAC_OUT[key])
        ws.delete_rows(r, 1)
        n += 1
        print(f"  Pending -> Rejected: {row.get('Date')} {key}")
    return n


def enrich_salmon(wb) -> int:
    """E — fill the FDA row the register has been sitting on."""
    ws = wb["Pending"]
    col = _cols(ws)
    n = 0
    for row in ws.iter_rows(min_row=2):
        if SALMON_URL not in str(row[col["URL"]].value or "").lower():
            continue
        if row[col["Pathogen"]].value == SALMON["Pathogen"]:
            continue
        for field, value in SALMON.items():
            if field in col:
                row[col[field]].value = value
        row[col["Notes"]].value = (str(row[col["Notes"]].value or "")
                                   + SALMON_NOTE).strip()
        n += 1
        print("  Pending enriched: Krasnyi Oktyabr Norwegian Salmon Sliced "
              "(FDA 2026-10-16)")
    return n


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    n_rm = remove_edsviken(wb)
    n_nm = fix_name_fields(wb)
    n_sp = annotate_superseded(wb)
    n_nf = reject_nafdac(wb)
    n_sa = enrich_salmon(wb)
    wb.save(XLSX)
    wb.close()
    print(f"name fixes: {n_nm} · SUPERSEDED reasons: {n_sp} · "
          f"NAFDAC rejected: {n_nf} · salmon enriched: {n_sa}")
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print(f"ROWS_REMOVED={n_rm}")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
