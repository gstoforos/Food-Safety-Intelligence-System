# -*- coding: utf-8 -*-
"""Operator pass 2026-10-10 (b): remove a non-food row from the register.

ROWS ARE REMOVED FROM Recalls — the commit message carries "remove-rows".

CDC MMWR 75(39), 8 Oct 2026, "Notes from the Field: Investigation of
Unapproved Botulinum Toxin Product Administered at a Medical Spa —
Colorado, 2025–2026" (https://www.cdc.gov/mmwr/volumes/75/wr/mm7539a2.htm)
was published to Recalls as a Tier-1 Clostridium botulinum outbreak. Read on
its own page 2026-10-10: three women developed symptoms after injections of a
non-FDA-approved botulinum toxin product at a Colorado medical spa. It is a
cosmetic injection, not food. Operator 2026-10-10: "It's not food, so we
don't have it — remove it from our data."

The row's fields were also malformed (Company and Brand held the MMWR
headline, Product "Colorado, 2025–2026").

What this does:
  * Recalls: the row is removed and archived to Rejected with a terminal
    reason, so the audit trail stays readable.
  * Weekly_Review: the Sunday review slice's copy is removed too (it is a
    live sheet; leaving it would put the row back in the review).
  * The publish gate gains a non-food-exposure rule in the same upload
    (pipeline/_publish_gate.py, tests/test_an_injection_is_not_food.py), so
    the next medical-spa or cosmetic-toxin report is refused at the gate.
"""
from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
TODAY = dt.date(2026, 10, 10).isoformat()

CDC_URL_KEY = "mm7539a2"
REASON = ("not_food: CDC MMWR 75(39) reports a non-FDA-approved botulinum toxin product "
          "INJECTED at a Colorado medical spa (three patients, all recovered) — a "
          "cosmetic/clinical exposure, not a food. Removed from Recalls by operator "
          "ruling 2026-10-10.")
BY = "operator review " + TODAY


def _head(ws):
    return [str(c.value) if c.value is not None else "" for c in ws[1]]


def _archive(wb, row, reason):
    arch = wb["Rejected"]
    out = []
    for h in _head(arch):
        if h in ("RejectReason", "RejectionReason"):
            out.append(reason)
        elif h == "RejectedBy":
            out.append(BY)
        elif h == "RejectedAt":
            out.append(TODAY)
        elif h == "Status":
            out.append("rejected")
        else:
            out.append(row.get(h, ""))
    arch.append(out)


def remove_from(wb, sheet, archive: bool) -> int:
    if sheet not in wb.sheetnames:
        return 0
    ws = wb[sheet]
    head = _head(ws)
    if "URL" not in head:
        return 0
    ucol = head.index("URL") + 1
    hits = [r for r in range(2, ws.max_row + 1)
            if CDC_URL_KEY in str(ws.cell(r, ucol).value or "")]
    for r in reversed(hits):
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(head)}
        print(f"  {sheet} row {r}: {row.get('Date')} {row.get('Source')} "
              f"{str(row.get('Company'))[:60]!r} — removing (not food)")
        if archive:
            _archive(wb, row, REASON)
        ws.delete_rows(r, 1)
    return len(hits)


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    n_rec = remove_from(wb, "Recalls", archive=True)
    n_rev = remove_from(wb, "Weekly_Review", archive=False)
    if n_rec or n_rev:
        wb.save(XLSX)
        from pipeline.merge_master import mirror_json_from_xlsx
        print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print(f"Weekly_Review copies removed: {n_rev}")
    print(f"ROWS_REMOVED={n_rec}")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
