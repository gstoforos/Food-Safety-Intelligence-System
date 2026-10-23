# -*- coding: utf-8 -*-
"""Review of the 2026-10-23 morning fix: three AESAN Pending rows queued by the
Exa daily-recall fallback at 07:41 UTC that held
tests/test_scraper_output_health.py::test_no_source_has_a_future_last_row RED.
NOTHING IS REMOVED FROM Recalls; the rows leave the Pending queue and are
archived to Rejected with the reason stated (archives are append-only).

1. https://www.aesan.gob.es/actualidad/actualidad-noticias — AESAN's NEWS
   LISTING page, not a recall notice. Company "Noticias", Product = the
   agency's own name, Date 2026-10-24 taken from an upcoming event on the
   listing ("24 Octubre 2026 ... Annual Workshop"). A future Date makes the
   AESAN collector look fresh.
2. https://aesan.gob.es/en/alertas/2026_72 and
   https://www.aesan.gob.es/eu/alertas/2026_72 — the English and Basque
   language variants of alert 2026_72 (Marvalmar cooked mussels, Listeria),
   which is already PUBLISHED on https://aesan.gob.es/alertas/2026_72
   (operator review 2026-10-23). Duplicates of a published recall.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"
JSON = ROOT / "docs" / "data" / "recalls.json"
TODAY = "2026-10-23"

TARGETS = {
    "aesan.gob.es/actualidad/actualidad-noticias":
        ("REJECTED — not a recall notice: AESAN's news listing page "
         "(Company 'Noticias'); its Date 2026-10-24 was read off an upcoming "
         "event on the listing. Operator review 2026-10-23."),
    "aesan.gob.es/en/alertas/2026_72":
        ("REJECTED — duplicate: English-language variant of AESAN alert "
         "2026_72, already published on https://aesan.gob.es/alertas/2026_72. "
         "Operator review 2026-10-23."),
    "aesan.gob.es/eu/alertas/2026_72":
        ("REJECTED — duplicate: Basque-language variant of AESAN alert "
         "2026_72, already published on https://aesan.gob.es/alertas/2026_72. "
         "Operator review 2026-10-23."),
}


def _head(ws):
    return [c.value for c in ws[1]]


def _archive(arch, row, reason):
    out = []
    for h in _head(arch):
        if h in ("RejectReason", "RejectionReason"):
            out.append(reason)
        elif h == "RejectedBy":
            out.append("operator review " + TODAY)
        elif h == "RejectedAt":
            out.append(TODAY)
        elif h == "Status":
            out.append("rejected")
        else:
            out.append(row.get(h, ""))
    arch.append(out)


def main() -> int:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    ws, arch = wb["Pending"], wb["Rejected"]
    head = _head(ws)
    ucol = head.index("URL") + 1
    n = 0
    for r in range(ws.max_row, 1, -1):
        url = str(ws.cell(r, ucol).value or "").strip().lower().rstrip("/")
        key = next((k for k in TARGETS if url.endswith(k)), None)
        if key is None:
            continue
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(head)}
        _archive(arch, row, TARGETS[key])
        ws.delete_rows(r, 1)
        n += 1
        print(f"  Pending -> Rejected: {url}")
    wb.save(XLSX)
    wb.close()
    print(f"AESAN Pending rows archived: {n}")
    from pipeline.merge_master import mirror_json_from_xlsx
    print(f"recalls.json re-mirrored: {mirror_json_from_xlsx(XLSX, JSON)} rows")
    print("ROWS_REMOVED=0")
    return 0


if __name__ == "__main__":                                  # pragma: no cover
    raise SystemExit(main())
