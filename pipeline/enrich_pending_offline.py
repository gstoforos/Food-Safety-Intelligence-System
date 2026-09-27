"""Fill Pending's missing Pathogen from the row's OWN Reason — offline.

WHY THIS EXISTS
---------------
On 2026-09-27 the whole Pending queue was run through the publish gate to see
what could be promoted. Of 50 rows, **one** passed. The dominant blocker was
not review and not a judgement call:

    18 rows  Pathogen is empty
    12 rows  Company is empty

and in most of those rows **the hazard was written in the row's own Reason
field the whole time** — "présence identifiée d'ochratoxine a", "Presence of
Salmonella", "elevated levels of lead", "presence of stones". Nothing needed
fetching, no model needed asking. The same shape as the RASFF notifId that sat
in Notes for six weeks while a test called it unrecoverable.

So this module copies the hazard the row already states into the field the gate
requires, deterministically, with no network and no LLM, so it can run on a
schedule twice a day instead of waiting for an operator.

THE RULES IT WILL NOT BREAK
---------------------------
1. It reads ONLY the row's own Reason (and Product as a fallback for the
   CFIA/FSAI listing rows, whose Reason IS the notice headline). It never
   fetches, never guesses and never consults a model.
2. It writes ONLY Pathogen. It will NOT fill Company, and that is deliberate:
   for RappelConso rows the only company-ish text on the row is the DISTRIBUTOR
   list in Notes ("magasins u", "leclerc thouars"), not the recalling firm.
   Promoting a supermarket as the recalling firm would be a fabricated
   attribution — and the register already refuses it: one such row carries
   "[confirm-agent hold: no recalling firm on the row (Brand 'Unbranded')]".
   Company stays empty and those rows stay in Pending, correctly.
3. Every value it writes is a canonical string ALREADY USED by published rows.
   The vocabulary below was derived by reading the register, not invented:
   "Foreign material (stones)", "Ochratoxin A", "Lead (heavy metal)" and the
   organism names are all spellings that 1,793 published rows already carry.
   0 of those 1,793 rows has an empty Pathogen, so this field is mandatory by
   convention as well as by the gate.
4. It stamps every change into Notes so the edit is auditable from the sheet.
5. Allergen-only hazards are NOT filled. Undeclared allergens are out of AFTS
   scope (policy 2026-07-29), so a row whose only stated hazard is an allergen
   is reported for rejection rather than quietly enriched into publishability.

USAGE
-----
    python -m pipeline.enrich_pending_offline              # report only
    python -m pipeline.enrich_pending_offline --apply      # write
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"

#: (regex, canonical Pathogen). Order matters — first match wins, so specific
#: patterns precede general ones. Every canonical value on the right-hand side
#: is a spelling already present in the published register.
RULES: Tuple[Tuple[str, str], ...] = (
    # ── mycotoxins and chemical contaminants ─────────────────────────────
    (r"ochratox|ochratoxine",                    "Ochratoxin A"),
    (r"aflatox",                                 "Aflatoxin"),
    (r"\bpatulin",                               "Patulin"),
    (r"\blead\b|\bplomb\b",                      "Lead (heavy metal)"),
    (r"\bcadmium\b",                             "Cadmium (heavy metal)"),
    (r"\bmercur|\bmercure\b",                    "Mercury (heavy metal)"),
    (r"\barsenic\b",                             "Arsenic (heavy metal)"),
    # ── bacteria and viruses ─────────────────────────────────────────────
    (r"listeria monocytogenes|list[ée]ria monocytogen",
                                                 "Listeria monocytogenes"),
    (r"\blisteri|\blist[ée]ri",                  "Listeria monocytogenes"),
    (r"e\.?\s?coli\s+stec|stec\b|verotoxin|shiga",
                                                 "Shiga toxin-producing E. coli (STEC)"),
    (r"escherichia coli|e\.?\s?coli",            "E. coli"),
    (r"salmonell",                               "Salmonella"),
    (r"campylobact",                             "Campylobacter"),
    (r"bacillus cereus|cereulid",                "Bacillus cereus"),
    (r"clostridium botulinum|botulis",           "Clostridium botulinum"),
    (r"clostridium perfringens",                 "Clostridium perfringens"),
    (r"staphylococc",                            "Staphylococcus aureus"),
    (r"cronobacter",                             "Cronobacter sakazakii"),
    (r"vibrio",                                  "Vibrio"),
    (r"yersinia",                                "Yersinia enterocolitica"),
    (r"hepatitis a|h[ée]patite a",               "Hepatitis A virus"),
    (r"norovirus|norwalk",                       "Norovirus"),
    (r"\bmould\b|\bmold\b|moisissure",           "Mold"),
    # ── physical / foreign material, with the material named where stated ─
    (r"stone|noyau|pierre",                      "Foreign material (stones)"),
    (r"\bglass\b|\bverre\b",                     "Foreign material (glass)"),
    (r"m[ée]tall?iqu|\bmetal\b",                 "Foreign material (metal)"),
    (r"\bplastic\b|plastique",                   "Foreign material (plastic)"),
    (r"\bwood\b|\bbois\b",                       "Foreign material (wood)"),
    (r"\binsect|\bpest\b|nuisible",              "Foreign material (pest)"),
    (r"corps [ée]tranger|foreign (?:body|bodies|object|material)",
                                                 "Foreign material"),
    # ── packaging integrity: a real, published hazard class ──────────────
    (r"herm[ée]ticit|[ée]tanch[ée]it|seal (?:defect|failure)|"
     r"perte d[’']?[ée]tanch|defect.{0,20}canette|swollen|bombage",
                                                 "Physical/foreign-body contamination"),
)

#: Hazards that are out of AFTS scope — never enriched into publishability.
ALLERGEN_ONLY = re.compile(
    r"undeclared|non[- ]d[ée]clar|allerg|peanut|arachide|gluten|sulphite|"
    r"sulfite|milk not declared|soy not declared", re.I)

#: Rows whose Reason is the notice HEADLINE rather than a hazard sentence.
#: For these the headline is the only hazard text the row has, so it is the
#: field to read. Identified by the scraper's own Notes stamp.
HEADLINE_SOURCES = ("CFIA HTML listing fallback", "FSAI HTML listing fallback",
                    "FDA HTML fallback")


def hazard_text(row: Dict) -> str:
    """The text this row states its hazard in — its own fields only."""
    parts = [str(row.get("Reason") or "")]
    notes = str(row.get("Notes") or "")
    if any(m in notes for m in HEADLINE_SOURCES):
        # The headline carries the hazard for these; Reason repeats it.
        parts.append(str(row.get("Product") or ""))
    return " ".join(parts)


def infer_pathogen(row: Dict) -> Tuple[Optional[str], str]:
    """(canonical Pathogen, note) inferred from the row's own text.

    Returns (None, reason) when nothing can be inferred without guessing.
    """
    text = hazard_text(row)
    if not text.strip():
        return None, "no hazard text on the row"

    low = text.lower()
    for pattern, canonical in RULES:
        m = re.search(pattern, low, re.I)
        if m:
            if ALLERGEN_ONLY.search(text) and "Foreign" not in canonical:
                # an allergen mentioned alongside a real hazard is fine;
                # only refuse when the allergen IS the hazard
                pass
            return canonical, f"matched {m.group(0)!r}"
    if ALLERGEN_ONLY.search(text):
        return None, ("hazard is an undeclared allergen — out of AFTS scope "
                      "(policy 2026-07-29), not enriched")
    return None, "no hazard term in the row's own text"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(XLSX))
    ap.add_argument("--sheet", default="Pending")
    ap.add_argument("--apply", action="store_true",
                    help="write the inferred Pathogen values into the sheet")
    args = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb[args.sheet]
    head = [str(c.value) for c in ws[1]]
    pcol = head.index("Pathogen") + 1
    ncol = head.index("Notes") + 1

    filled: List[str] = []
    skipped: List[str] = []
    for r in range(2, ws.max_row + 1):
        row = {head[i]: ws.cell(r, i + 1).value for i in range(len(head))}
        if str(row.get("Pathogen") or "").strip():
            continue
        got, why = infer_pathogen(row)
        label = f"{str(row.get('Source'))[:14]:15} {str(row.get('Product'))[:44]}"
        if not got:
            skipped.append(f"{label}  — {why}")
            continue
        filled.append(f"{label}  -> {got}   ({why})")
        if args.apply:
            ws.cell(r, pcol).value = got
            stamp = (f"[offline-enrich {dt.date.today().isoformat()}: Pathogen "
                     f"'' -> {got!r}, read from this row's own hazard text; "
                     f"no fetch, no model]")
            prev = str(ws.cell(r, ncol).value or "")
            if stamp not in prev:
                ws.cell(r, ncol).value = (prev + " " + stamp).strip()

    print(f"FILLED {len(filled)}:")
    for f in filled:
        print("   " + f)
    print(f"\nLEFT ALONE {len(skipped)}:")
    for s in skipped:
        print("   " + s)

    if args.apply and filled:
        wb.save(args.xlsx)
        print(f"\nwrote {len(filled)} Pathogen value(s) to {args.xlsx}")
        print("NOTE: Company is never filled here — see rule 2 in the "
              "module docstring.")
    elif not args.apply:
        print("\n(report only — pass --apply to write)")
    return 0


if __name__ == "__main__":                                 # pragma: no cover
    raise SystemExit(main())
