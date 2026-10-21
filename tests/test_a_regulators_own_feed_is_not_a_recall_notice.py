# -*- coding: utf-8 -*-
"""A regulator's syndication feed is a document listing many recalls.

THE ROW THAT FOUND THIS (main at 0ae22b9c, 2026-10-21). Recalls carried,
DateAdded 2026-10-21, Source "CFIA", Date 2026-10-14, Tier 1, Listeria
monocytogenes:

    Company  "Health and safety"
    Brand    "Health and safety"
    Product  "and alerts"
    URL      https://recalls-rappels.canada.ca/en/feed/cfia-alerts-recalls

Those three field values are the feed's own channel title, "Health and
safety recalls and alerts", split at the word "recalls". The recall the
feed happened to be carrying was already published on its own URL
(/en/alert-recall/various-brands-organic-cheeses-recalled-due-listeria-
monocytogenes), so the row was a second, nameless copy pointing a
subscriber at an XML document.

Every existing URL rule passed it: the path is not in _LANDING_PATHS, it is
not a CMS index prefix, the host is the regulator's own host, and the path
is long and slug-shaped.

The rule is scoped to REGULATOR hosts on purpose — see the comment on
url_path_is_a_feed in pipeline/_publish_gate.py. One published row cites a
news feed under a deliberate operator exception (NFSA (EG), 2026-10-11,
"News-authority source: see.news (no official NFSA per-recall portal)") and
this rule must not refuse it as a side effect.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from pipeline._publish_gate import (  # noqa: E402
    publish_blockers,
    url_path_is_a_feed,
)

CFIA_FEED = "https://recalls-rappels.canada.ca/en/feed/cfia-alerts-recalls"


def _row(url: str, **over) -> dict:
    row = {
        "Date": "2026-10-14",
        "Source": "CFIA",
        "Company": "Fromagerie L'Ancêtre Inc.",
        "Brand": "Farm Boy",
        "Product": "Farm Boy Organic Swiss Cheese 190 g",
        "Pathogen": "Listeria monocytogenes",
        "Reason": "Presence of Listeria monocytogenes",
        "Class": "Recall",
        "Country": "Canada",
        "Region": "North America",
        "Tier": 1,
        "Outbreak": 0,
        "URL": url,
        "Notes": "",
    }
    row.update(over)
    return row


# ── the path test, in isolation ────────────────────────────────────────────

@pytest.mark.parametrize("path", [
    "/feed/cfia-alerts-recalls",
    "/en/feed/cfia-alerts-recalls",
    "/fr/fil-de-nouvelles/acia-rappels-avis-securite",
    "/feeds/food-recalls",
    "/rss",
    "/rss/articles/abc123",
    "/rss-feeds",
    "/atom",
    "/news/rss.xml",
    "/alerts/feed.xml",
    "/warnungen.rss",
])
def test_a_feed_path_is_recognised(path: str) -> None:
    assert url_path_is_a_feed(path), path


@pytest.mark.parametrize("path", [
    "",
    "/en/alert-recall/various-brands-organic-cheeses-recalled-due-listeria-"
    "monocytogenes",
    # "feed" and "rss" inside a slug are not a feed. The gate's own stated
    # philosophy: a rule that cries wolf on good rows gets switched off.
    "/recalls-alerts/animal-feed-recalled-due-to-salmonella",
    "/food-recalls/recall-alert/premium-chicken-feed-pellets-20kg",
    "/alert-recall/feedlot-beef-recalled-due-to-e-coli-o157h7",
    "/fiche-rappel/23876/interne",
    "/press-release/rssmann-dairy-listeria",
], ids=lambda p: (p or "empty")[:48])
def test_a_slug_containing_feed_or_rss_is_not_a_feed(path: str) -> None:
    assert not url_path_is_a_feed(path), path


# ── the gate ───────────────────────────────────────────────────────────────

def test_the_gate_refuses_a_regulator_feed_url() -> None:
    problems = publish_blockers(_row(CFIA_FEED))
    assert any("syndication feed" in p for p in problems), problems


def test_the_gate_still_accepts_the_real_notice() -> None:
    url = ("https://recalls-rappels.canada.ca/en/alert-recall/various-brands-"
           "organic-cheeses-recalled-due-listeria-monocytogenes")
    assert publish_blockers(_row(url)) == []


def test_the_news_authority_exception_is_not_collateral_damage() -> None:
    """A news feed URL is a DIFFERENT rule's business.

    NFSA (EG) 2026-10-11 is published on a news.google.com/rss/articles/…
    URL under an operator exception recorded in its own Notes. Whatever
    else refuses or allows that row, it must not be THIS rule, which is
    about a regulator's own feed.
    """
    url = ("https://news.google.com/rss/articles/"
           "CBMifkFVX3lxTFBUWmtsaG4ySWVMZTFaQzBSVUIxSEM?oc=5")
    row = _row(url, Source="NFSA (EG)", Country="Egypt", Region="Africa",
               Company="Yasensvit Ltd", Brand="", Product="Eggs",
               Pathogen="Salmonella Enteritidis",
               Notes=("News-authority source: see.news (no official NFSA "
                      "per-recall portal)"))
    assert not any("syndication feed" in p
                   for p in publish_blockers(row))


# ── the register ───────────────────────────────────────────────────────────

def test_no_published_row_cites_a_regulators_own_feed() -> None:
    import openpyxl

    xlsx = ROOT / "docs" / "data" / "recalls.xlsx"
    if not xlsx.exists():                                    # pragma: no cover
        pytest.skip("workbook not present")
    from pipeline.regulatory_domains import is_regulator_url

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    offenders = []
    try:
        for sheet in ("Recalls", "Pending", "Weekly_Review"):
            if sheet not in wb.sheetnames:
                continue
            rows = wb[sheet].iter_rows(values_only=True)
            head = list(next(rows))
            if "URL" not in head:
                continue
            ui = head.index("URL")
            for row in rows:
                url = str((row[ui] if row else "") or "").strip()
                if not url.startswith(("http://", "https://")):
                    continue
                tail = url.split("?", 1)[0].split("#", 1)[0]
                after_host = tail.split("://", 1)[-1]
                path = (("/" + after_host.split("/", 1)[1])
                        if "/" in after_host else "")
                if url_path_is_a_feed(path.rstrip("/").lower()) \
                        and is_regulator_url(url):
                    offenders.append((sheet, url[:100]))
    finally:
        wb.close()
    assert not offenders, (
        "row(s) cite a regulator's syndication feed instead of a per-recall "
        f"notice: {offenders}")
