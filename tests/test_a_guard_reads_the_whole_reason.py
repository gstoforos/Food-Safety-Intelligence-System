"""An annotation on an archived reason must not blind the guards that read it.

WHY THIS EXISTS (2026-10-04)
============================
merge_master.load_rejected_urls built the value every re-promotion exception
tests, and built it as

    desc = f"{by}: {why[:160]} | {notes[:400]}"

The slices were there to keep a log line short. But that value is also what
the exceptions SEARCH, and every one of them is a substring search:

    REPAIRABLE_DEFECTS      "no matching hazard category", "pathogen is
                            empty", "company and brand are the same", ...
    _url_guard.reject_refusal()
    _reversal_excuses()
    _is_terminal_rejection()

So any text PREPENDED to RejectionReason pushed the real defect name past
character 160 and all four went blind.

The register does this to itself. promote_gate_passing.
supersede_archived_copies prepends a ~250-character

    "[SUPERSEDED <date> — the defect named below was repaired and this
     recall is PUBLISHED in Recalls. ...]"

marker. The marker alone overflows the window, so every row the promoter
has ever marked SUPERSEDED became permanently unre-promotable — the exact
dead end the first exception was written to remove, reintroduced by the fix
for a different defect. 38 archive rows carried such a marker on
2026-10-04, and 274 carried a reason longer than the window.

THE ROW THAT FOUND IT. Italy, Franchi Salumi, salamella dolce, archived
"unknown: No matching hazard category — defer to manual review" on the
FOREIGN_MATTER vocabulary gap fixed the same morning. With the gap fixed,
the row repaired and the archive copy marked SUPERSEDED, it still would not
promote, and the only thing the log said was "re-promotion BLOCKED".
"""
from __future__ import annotations

from pipeline.merge_master import load_rejected_urls, _normalize_url_for_dedup


REPAIRABLE = "no matching hazard category"
MARKER = (
    "[SUPERSEDED 2026-10-04 — the defect named below was repaired and this "
    "recall is PUBLISHED in Recalls. This row is kept only as the audit "
    "trail of the original refusal; the Recalls copy is what the register "
    "says. Padding to the length the real marker reaches so this test fails "
    "if the window ever comes back.] ")


def _book(tmp_path, reason):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Recalls"
    ws.append(["URL"])
    wr = wb.create_sheet("Weekly_Rejected")
    wr.append(["URL", "RejectedBy", "RejectionReason", "Notes", "Reviewed"])
    wr.append(["https://www.example.gov/alert/1", "gap_finder/it/rules.py",
               reason, "", "SUPERSEDED"])
    p = tmp_path / "recalls.xlsx"
    wb.save(p)
    return p


def test_the_defect_name_survives_a_long_annotation(tmp_path):
    assert len(MARKER) > 160, "the marker must overflow the old window"
    p = _book(tmp_path, MARKER + "unknown: " + REPAIRABLE.title()
              + " — defer to manual review.")
    got = load_rejected_urls(p)
    key = _normalize_url_for_dedup("https://www.example.gov/alert/1")
    assert key in got
    assert REPAIRABLE in got[key].lower(), (
        "the repairable-defect name was cut off by a prefix annotation, so "
        "merge_master's first re-promotion exception can never fire: "
        + got[key][:200])


def test_a_notes_verdict_past_400_chars_still_blocks(tmp_path):
    """Notes was sliced too, and the terminal verdict often lives there."""
    verdict = "REJECTED: pet_food_out_of_scope — AFTS-FSIS is a HUMAN-food register."
    p = _book(tmp_path, "unknown")
    import openpyxl
    wb = openpyxl.load_workbook(p)
    wb["Weekly_Rejected"].cell(2, 4).value = ("x" * 500) + " " + verdict
    wb.save(p)
    got = load_rejected_urls(p)
    key = _normalize_url_for_dedup("https://www.example.gov/alert/1")
    assert "pet_food_out_of_scope" in got[key], (
        "a terminal verdict recorded past character 400 of Notes was "
        "invisible to the terminal-rejection guard")
