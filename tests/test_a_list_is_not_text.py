"""A list is not text (2026-10-01): Reason "['a', 'a']" reached Recalls."""
from __future__ import annotations

from pathlib import Path

import pytest

from pipeline._list_text import as_text

ROOT = Path(__file__).resolve().parents[1]
REPR = "['Contamination of product with salmonella', 'Contamination of product with salmonella']"


@pytest.mark.parametrize("val,want", [
    (REPR, "Contamination of product with salmonella"),
    (["a", "b", "A", ""], "a; b"),
    ('["x", "y"]', "x; y"),
    ("[Updated] Listeria recall", "[Updated] Listeria recall"),
    ("plain", "plain"),
    ("[1, 2]", "[1, 2]"),
])
def test_as_text(val, want):
    assert as_text(val) == want


@pytest.mark.parametrize("mod", [
    "pipeline.extractor", "pipeline.official_feeds.extractor",
    "pipeline.gap_finder.extractor", "pipeline.gap_finder_gr.extractor",
])
def test_every_extractor_flattens_a_model_list(mod):
    import importlib
    m = importlib.import_module(mod)
    assert m._safe({"reason_en": ["a", "a"]}, "reason_en") == "a"


def test_the_writer_flattens_a_list_repr():
    from openpyxl import Workbook
    from pipeline.merge_master import _write_sheet, SCHEMA
    row = {k: "" for k in SCHEMA}
    row.update({"Date": "2026-09-25", "URL": "https://x.example/a",
                "Reason": REPR, "Pathogen": "Salmonella"})
    wb = Workbook()
    _write_sheet(wb, "Pending", SCHEMA, [row])
    col = SCHEMA.index("Reason") + 1
    assert wb["Pending"].cell(2, col).value == "Contamination of product with salmonella"


def test_the_register_holds_no_list_repr():
    import openpyxl
    wb = openpyxl.load_workbook(ROOT / "docs" / "data" / "recalls.xlsx", read_only=True)
    bad = []
    for name in ("Recalls", "Weekly_Review", "Pending"):
        if name not in wb.sheetnames:
            continue
        it = wb[name].iter_rows(values_only=True)
        head = next(it)
        cols = [head.index(c) for c in ("Reason", "Product", "Company") if c in head]
        for r in it:
            for i in cols:
                v = r[i]
                if isinstance(v, str) and as_text(v) != v:
                    bad.append((name, v[:60]))
    assert not bad, bad
