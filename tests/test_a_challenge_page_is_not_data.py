"""A bot-wall / challenge page is never data (2026-09-30).

salute.gov.it answers datacentre fetches with a Gcore challenge page and
HTTP 200. Its <title> became the Product of eleven Italian rows ("Gcore").
The fix sits at every point a wall can enter, and this file sweeps them:
fetch (article fetcher, provenance, both reviewers' page readers, and the
heredoc copies that actually run), title extraction, the gap-finder
garbage guard, the url-guard's reachability vocabulary, and the writer
choke point.
"""
from __future__ import annotations

import re
from pathlib import Path
from unittest import mock

import pytest

from pipeline import _bot_wall as bw

ROOT = Path(__file__).resolve().parents[1]

GCORE = "<html><head><title>Gcore</title></head><body>Please wait</body></html>"
CF = ('<html><head><title>Just a moment...</title></head><body>'
      '<div id="cf-chl-widget"></div></body></html>')
REAL_WITH_CF_SCRIPT = (
    '<html><head><title>Richiamo hamburger di scottona</title>'
    '<script src="/cdn-cgi/challenge-platform/scripts/x.js"></script></head>'
    '<body>' + "Il Ministero della Salute comunica il richiamo. " * 60 +
    '</body></html>')


@pytest.mark.parametrize("title", ["Gcore", "gcore", " Just a moment... ",
                                   "Attention Required! | Cloudflare",
                                   "Access Denied", "403 Forbidden"])
def test_wall_titles(title):
    assert bw.is_wall_title(title)


@pytest.mark.parametrize("title", ["Gcore protein bar", "Hamburger di scottona",
                                   "Access Denied brand crisps", "", None])
def test_real_titles_are_not_walls(title):
    assert not bw.is_wall_title(title)


def test_wall_html():
    assert bw.is_bot_wall_html(GCORE)
    assert bw.is_bot_wall_html(CF)


def test_a_real_page_that_loads_a_challenge_script_is_not_a_wall():
    assert not bw.is_bot_wall_html(REAL_WITH_CF_SCRIPT)


def test_screen_only_touches_ok_walls():
    assert bw.screen("u", GCORE, "ok") == ("u", "", "bot_wall")
    assert bw.screen("u", REAL_WITH_CF_SCRIPT, "ok")[2] == "ok"
    assert bw.screen("u", "", "http_403") == ("u", "", "http_403")


# ── fetch ────────────────────────────────────────────────────────────────

def test_the_gap_finder_fetcher_screens():
    import pipeline.gap_finder.article_fetcher as af
    with mock.patch.object(af, "_fetch_html_unscreened",
                           return_value=("u", GCORE, "ok")):
        assert af.fetch_html("u", None) == ("u", "", "bot_wall")
    assert af.extract_title(GCORE) == ""
    assert af.extract_title("<h1>Richiamo hamburger scottona</h1>") == \
        "Richiamo hamburger scottona"


def test_provenance_treats_a_wall_as_infrastructure_not_content():
    from pipeline import _provenance as pv
    resp = mock.Mock(status_code=200, text=GCORE + "x" * 0)
    with mock.patch.object(pv, "_fetch_response", return_value=(resp, "")):
        assert pv.fetch_text("https://www.salute.gov.it/x") == ("", "bot_wall")
    assert not pv.is_dead_status("bot_wall"), (
        "a wall must never read as 'the page does not exist'")


READERS = [
    "pipeline/recall_url_agent.py",
    "pipeline/recall_review_agent.py",
    ".github/workflows/recall-url-agent.yml",
    ".github/workflows/recall-review-agent.yml",
    ".github/workflows/recallreviewagent.yml",
]


@pytest.mark.parametrize("rel", READERS)
def test_every_reviewer_page_reader_screens(rel):
    text = (ROOT / rel).read_text(encoding="utf-8")
    assert "def _fetch_page_text" in text
    body = text[text.index("def _fetch_page_text"):]
    body = body[:body.index("\ndef ", 1) if "\ndef " in body[1:] else None]
    assert "is_bot_wall_html" in body and '"bot_wall"' in body, (
        f"{rel}: the reviewer's page reader passes a challenge page to the "
        f"model as if it were the notice")


def test_reviewer_reader_behaviour():
    from pipeline import recall_url_agent as a1
    fake = mock.Mock(status_code=200, text=GCORE)
    cf = mock.Mock(get=mock.Mock(return_value=fake))
    with mock.patch.dict("sys.modules", {"curl_cffi": mock.Mock(requests=cf),
                                         "curl_cffi.requests": cf}):
        assert a1._fetch_page_text("https://www.salute.gov.it/x") == \
            ("", "bot_wall")


def test_the_guard_reads_a_wall_as_reachability():
    from pipeline._url_guard import reject_refusal
    row = {"URL": "https://www.salute.gov.it/new/sites/default/files/"
                  "external_data/avvisi_sicurezza_alimentare/x.pdf"}
    for why in ("fetch returned status bot_wall", "Gcore challenge page",
                "blocked by a CAPTCHA", "browser validation wall"):
        assert reject_refusal(row, why), why
    assert not reject_refusal(row, "pet food, out of scope")


# ── guards and writer ────────────────────────────────────────────────────

def test_the_gap_finder_guard_calls_a_wall_title_garbage():
    from pipeline._gap_finder_guards import product_is_garbage
    assert product_is_garbage("Gcore")
    assert not product_is_garbage("Hamburger di scottona 200 g")


def test_the_writer_blanks_wall_titles_on_every_sheet(tmp_path):
    from pipeline.merge_master import save_xlsx_with_pending
    import openpyxl
    row = {"Date": "2026-09-25", "Source": "Salute", "Company": "Ambrosini Carni",
           "Brand": "Gran Selezione Ambrosini", "Product": "Gcore",
           "Pathogen": "Salmonella", "Reason": "Salmonella spp.",
           "Country": "Italy", "URL": "https://www.salute.gov.it/x.pdf",
           "Status": "pending", "Notes": ""}
    x = tmp_path / "r.xlsx"
    save_xlsx_with_pending([], [dict(row)], x)
    wb = openpyxl.load_workbook(x, read_only=True)
    ws = wb["Pending"]
    it = ws.iter_rows(values_only=True)
    h = next(it)
    got = dict(zip(h, next(it)))
    assert (got["Product"] or "") == ""
    assert got["Company"] == "Ambrosini Carni"
    assert "bot-wall title removed" in (got["Notes"] or "")


# ── the re-promotion guard (2026-09-30) ──────────────────────────────────

def _repairable():
    src = (ROOT / "pipeline" / "merge_master.py").read_text(encoding="utf-8")
    block = src[src.index("REPAIRABLE_DEFECTS = ("):]
    block = block[:block.index("\n            )")]
    lines = [l for l in block.splitlines() if not l.strip().startswith("#")]
    return re.findall(r'"([^"]+)"', "\n".join(lines))


def test_a_wall_archived_row_can_come_back_once_repaired():
    """Ambrosini (Salmonella) was archived as 'No matching hazard category'
    because its fields held a Gcore page; R.J. King (Staph aureus) as
    'Company and Brand are the same long headline string'. Field defects."""
    rep = _repairable()
    assert "no matching hazard category" in rep
    assert "company and brand are the same" in rep


@pytest.mark.parametrize("verdict", [
    "pathogen_out_of_scope: 'Peanut'", "allergen: Undeclared allergen — out of scope",
    "Duplicate URL already exists", "not_food", "pet_food_out_of_scope",
    "no-authority-url: no official NÉBIH press-release URL",
])
def test_no_content_verdict_is_repairable(verdict):
    assert not any(d in verdict.lower() for d in _repairable()), verdict


def test_every_repairable_defect_also_supersedes_its_archive_copy(tmp_path):
    """merge_master lets a repaired row back in; promote_gate_passing then
    marks the archived copy SUPERSEDED. An entry in the first and not the
    second leaves a recall both published and rejected (CFIA R.J. King
    lobster, 2026-09-30).

    2026-10-05: this compared the two lists as SOURCE TEXT, which broke the
    moment promote_gate_passing replaced its three hand-listed
    "<field> is empty" phrases with a regex over publish-gate rule 4's own
    field list — a strictly WIDER rule that the string comparison read as
    three deletions. The contract was never "the same literals appear in
    both files"; it is "a row archived for a repairable defect gets
    superseded once the recall is published". So ask the function.
    """
    import openpyxl
    from pipeline.promote_gate_passing import supersede_every_published_url

    url = "https://example.gov/recall/one"
    missing = []
    for defect in sorted(_repairable()):
        wb = openpyxl.Workbook()
        r = wb.active
        r.title = "Recalls"
        r.append(["Date", "Product", "URL"])
        r.append(["2026-09-18", "A recall", url])
        a = wb.create_sheet("Rejected")
        a.append(["Date", "Product", "URL", "Notes", "Status", "RejectedBy",
                  "RejectReason"])
        a.append(["2026-09-20", "A recall", url, "", "rejected", "unknown",
                  f"reviewer 2: {defect}"])
        p = tmp_path / (re.sub(r"\W+", "_", defect) + ".xlsx")
        wb.save(p)
        if supersede_every_published_url(p) != 1:
            missing.append(defect)
    assert not missing, (
        f"merge_master.REPAIRABLE_DEFECTS lets a row archived for {missing} "
        f"back into Recalls, but promote_gate_passing does not retire the "
        f"archive copy, so the recall ends up both published and rejected.")


def test_a_reversed_policy_can_let_an_archived_row_back():
    """FSIS Sempio: archived 2026-09-29 as out_of_scope_import_reinspection,
    ruled in scope 2026-09-30. The re-promotion guard must consult the
    policy-reversal table, not only the repairable-defect list."""
    from pipeline.merge_master import _reversal_excuses
    old = ("operator review 2026-09-29: out_of_scope_import_reinspection — "
           "FSIS import re-inspection violations are deliberately excluded")
    assert _reversal_excuses(old)
    assert not _reversal_excuses(
        "operator review 2026-10-02: out_of_scope_import_reinspection")
    src = (ROOT / "pipeline" / "merge_master.py").read_text("utf-8")
    assert '_repromote_kind = "reversal"' in src
