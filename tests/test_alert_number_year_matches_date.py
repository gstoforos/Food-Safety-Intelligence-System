# -*- coding: utf-8 -*-
"""A notice numbered 027/2022 may not be published with a 2026 Date.

FOUND 2026-10-17. Two rows were published out of NAFDAC's discovery
channel with the day they were FOUND in Date instead of the day the
regulator issued the alert:

    Recalls  2026-10-16  NAFDAC (NG)  Multivitamin Supplements
             .../public-alert-no-023-2025-...     released 5 July 2025
    Recalls  2026-10-13  NAFDAC (NG)  Alcoholic Bitters
             .../public-alert-no-027-2022-...     released 10 May 2022

A 15-month and a 4½-year error, both of them landing in October 2026's
totals, the daily briefs and the open week's page. Neither needed a fetch
to catch: many regulators number a public alert "<serial>/<year>" and keep
that number in the URL slug, so the row's OWN URL contradicts its own
Date. The register's archives hold refused NAFDAC alerts numbered 2018,
2019, 2021, 2022, 2023, 2024 and 2025 — all stamped with their 2026
discovery date — so the channel produces this shape constantly and the
publish gate had nothing that looked for it.

This test is deliberately narrow. It fires only on a number that appears
in an alert/notice-number CONTEXT in the URL ("public-alert-no-027-2022",
"aviso-no-14-2025"), never on a bare year in a path or on a date-like
slug, because a slug year that merely repeats the publication year is
normal and harmless. A mismatch is always worth a human look: either the
Date is the scrape date, or the row re-publishes an old alert.
"""
from __future__ import annotations

import datetime as _dt
import re
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "data" / "recalls.xlsx"

#: "public-alert-no-023-2025", "alert-no_14/2025", "aviso-no-3-2024".
#: The serial comes first, the year second; both are required, and the
#: word that introduces them must be an alert/notice word.
ALERT_NO = re.compile(
    r"(?:public[-_]?alert|alert|alerta|aviso|notice|komunikat|ostrzezenie)"
    r"[-_a-z]*[-_]no?[-_]?(\d{1,4})[-/](20\d{2})",
    re.IGNORECASE,
)

#: Rows the operator has ruled on explicitly. A row only belongs here
#: with a reason and a date, and only when the Date is right and the
#: number is genuinely from another year (a re-issued or updated alert).
ALLOWED: dict[str, str] = {}

LIVE_SHEETS = ("Recalls", "Pending")


def _year(value) -> int | None:
    if isinstance(value, (_dt.datetime, _dt.date)):
        return value.year
    m = re.match(r"\s*(20\d{2})-\d{2}-\d{2}", str(value or ""))
    return int(m.group(1)) if m else None


def _rows(sheet: str):
    import openpyxl

    if not XLSX.exists():                                   # pragma: no cover
        pytest.skip("no workbook")
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:                      # pragma: no cover
            pytest.skip("%s not present" % sheet)
        it = wb[sheet].iter_rows(values_only=True)
        hdr = list(next(it))
        return [dict(zip(hdr, r)) for r in it
                if r and any(v is not None and str(v).strip() for v in r)]
    finally:
        wb.close()


@pytest.mark.parametrize("sheet", LIVE_SHEETS)
def test_no_row_publishes_an_alert_from_another_year(sheet):
    offenders = []
    for row in _rows(sheet):
        url = str(row.get("URL") or "").strip()
        if not url or url in ALLOWED:
            continue
        year = _year(row.get("Date"))
        if year is None:
            continue                       # an empty Date is a different gate
        m = ALERT_NO.search(url)
        if not m:
            continue
        alert_year = int(m.group(2))
        if alert_year != year:
            offenders.append((str(row.get("Date"))[:10], row.get("Source"),
                              m.group(0), url[:110]))

    assert not offenders, (
        "%d row(s) in %s carry an alert number from a different year than their "
        "Date:\n  %s\n\nThe regulator numbers its alerts <serial>/<year> and keeps "
        "that number in the URL, so the row's own URL contradicts its own Date. "
        "Either the Date is the day the row was discovered rather than the day the "
        "alert was issued, or the row re-publishes an old alert. Read the notice, "
        "set Date from the notice, and re-check the hazard against the scope rule "
        "before publishing it. If the mismatch is genuine (a re-issued alert), add "
        "the URL to ALLOWED in this file with the reason and the date."
        % (len(offenders), sheet, "\n  ".join(map(str, offenders))))


# ── the regex itself: narrow on purpose ────────────────────────────────────

@pytest.mark.parametrize("url,expected", [
    ("https://nafdac.gov.ng/public-alert-no-027-2022-adulterated-lots", "2022"),
    ("https://nafdac.gov.ng/public-alert-no-023-2025-recall-of-multi", "2025"),
    ("https://x.gov/alerta-no-14-2024-algo", "2024"),
    # a bare year in a path is NOT an alert number
    ("https://www.nvwa.nl/actueel/nieuws/2026/10/16/veiligheidswaarschuwing", None),
    ("https://www.lebensmittelwarnung.de/Meldungen/2026/10_Oktober/261016_18", None),
    ("https://rappel.conso.gouv.fr/fiche-rappel/23830/interne", None),
    ("https://www.fsai.ie/news-and-alerts/food-alerts/recall-of-a-batch", None),
    ("https://webgate.ec.europa.eu/rasff-window/screen/notification/846682", None),
])
def test_the_pattern_only_matches_an_alert_number(url, expected):
    m = ALERT_NO.search(url)
    assert (m.group(2) if m else None) == expected
