"""Fixes from the review of the 2026-10-01 night and 2026-10-02 morning
reviewer runs (operator 2026-10-02: "fix everything").

Each test names the row that exposed the defect.
"""
from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


# ── 1. Gap-finder rows: authority date first, and a Class ──────────────────

def test_gap_row_takes_the_authority_date_and_gets_a_class():
    """GIS Pilos protein pudding: page dated 2026-08-21, row dated 2026-09-21
    by the model; and Class was empty, so the publish gate refused the row
    on every run although reviewer 2 had approved it."""
    from pipeline.gap_finder.extractor import build_pending_row, notice_class

    class Cfg:
        authority_short, name_en, code = "GIS (PL)", "Poland", "pl"

    class Cls:
        tier, outbreak_qualifies = 1, False

    verified = {"efet_date_iso": "2026-08-21",
                "efet_title": "Ostrzeżenie publiczne dotyczące żywności: "
                              "Bacillus cereus w puddingu proteinowym",
                "efet_url": "https://www.gov.pl/web/gis/x",
                "news_source_domain": "gazeta.pl"}
    row = build_pending_row(verified, Cls(), {"date_iso": "2026-09-21",
                                               "company": "Lidl"}, Cfg())
    assert row["Date"] == "2026-08-21"
    assert row["Class"] == "Alert"
    assert notice_class("Media statement: Deli Hummus range recall") == "Recall"


def test_gap_row_falls_back_to_the_model_date_only_without_one():
    from pipeline.gap_finder.extractor import build_pending_row

    class Cfg:
        authority_short, name_en, code = "NCC (ZA)", "South Africa", "za"

    class Cls:
        tier, outbreak_qualifies = 1, False

    row = build_pending_row({"efet_date_iso": "", "efet_title": "x"}, Cls(),
                            {"date_iso": "2026-09-29"}, Cfg())
    assert row["Date"] == "2026-09-29"


# ── 2. Reviewer 3 waits for reviewer 1 ─────────────────────────────────────

def test_reviewer_3_leaves_a_row_reviewer_1_has_not_finished():
    from pipeline.recall_confirm_agent import _still_with_reviewer_1 as wait
    now = dt.datetime(2026, 10, 2, 3, 0, tzinfo=dt.timezone.utc)
    assert wait({"Status": "pending_gap_v1",
                 "ScrapedAt": "2026-10-01T10:00:00+00:00"}, now)
    assert wait({"Status": "Pending_Gap", "ScrapedAt": "2026-10-01T10:00:00Z"}, now)
    # bounded: after the grace period it is judged like any other row
    assert not wait({"Status": "pending_gap", "ScrapedAt": "2026-09-20T10:00:00"}, now)
    # reviewer 2's lane and rows with no age are judged now
    assert not wait({"Status": "pending", "ScrapedAt": "2026-10-01T10:00:00"}, now)
    assert not wait({"Status": "pending_gap"}, now)


def test_the_wait_is_wired_into_the_decision_loop():
    src = (ROOT / "pipeline" / "recall_confirm_agent.py").read_text(encoding="utf-8")
    assert "if probs and _still_with_reviewer_1(row):" in src


# ── 3. RejectedBy names the current reviewers ──────────────────────────────

@pytest.mark.parametrize("reason,who", [
    ("Confirmer: row was at pending and Date is empty", "reviewer 3 (confirm-agent)"),
    ("URL agent: No specific official recall page found", "reviewer 1 (url-agent)"),
    ("Review agent: field integrity", "reviewer 2 (review-agent)"),
    ("Arrived already marked rejected; confirmer did not re-review",
     "earlier reviewer (archived by reviewer 3)"),
    ("claude-check: x", "claude-check"),
    ("something else", "unknown"),
])
def test_rejected_by_is_read_from_the_reason(reason, who):
    from pipeline.merge_master import _rejected_by_from_reason
    assert _rejected_by_from_reason(reason) == who


# ── 4. An "unknown:" label does not make a scope verdict retryable ─────────

def test_a_scope_verdict_by_an_unnamed_reviewer_stays_final():
    """RappelConso 23644 (taboulé, shelf-life extension) was archived at
    23:11 and re-ingested at 04:11: the RejectedBy label "unknown:" matched
    the transient marker "unknown"."""
    from pipeline.merge_master import _is_terminal_rejection as final
    desc = ("unknown: URL agent: The hazard is not a microbial pathogen, it is "
            "a prolongation of shelf life || unknown: Arrived already marked "
            "rejected; confirmer did not re-review")
    assert final(desc)
    # a reason that literally reads "unknown", or a fetch failure, stays retryable
    assert not final("unknown")
    assert not final("unknown: URL agent: http_error 404")


# ── 5. A row the other side archived does not come back in a merge ─────────

def _book(path, pending_urls, rejected_urls):
    from openpyxl import Workbook
    from pipeline.merge_master import PENDING_SCHEMA, RECALLS_SCHEMA
    wb = Workbook()
    ws = wb.active
    ws.title = "Recalls"
    ws.append(RECALLS_SCHEMA)
    r = {c: "" for c in RECALLS_SCHEMA}
    r.update(Date="2026-09-01", Source="FDA", Company="Kept Co", Product="p",
             URL="https://www.fda.gov/kept")
    ws.append([r[c] for c in RECALLS_SCHEMA])
    wp = wb.create_sheet("Pending")
    wp.append(PENDING_SCHEMA)
    for i, u in enumerate(pending_urls):
        row = {c: "" for c in PENDING_SCHEMA}
        row.update(Date="2026-10-01", Source="FSANZ (AU)", Company=f"Co{i}",
                   Product=f"Product {u[-1]}", URL=u, Status="pending")
        wp.append([row[c] for c in PENDING_SCHEMA])
    wj = wb.create_sheet("Weekly_Rejected")
    wj.append(["Date", "Source", "Company", "Product", "URL", "RejectionReason"])
    for u in rejected_urls:
        wj.append(["2026-10-01", "FSANZ (AU)", "x", f"Product {u[-1]}", u, "Confirmer: x"])
    wb.save(path)


def test_row_merge_does_not_resurrect_an_archived_row(tmp_path):
    """Night of 2026-10-01: reviewer 1 archived FSANZ Sunlife, NCC BM Foods
    and FDA H-1380-2026; reviewer 2 lost the push race and the row union put
    all three back into Pending."""
    import openpyxl
    from pipeline.xlsx_merge import merge_xlsx_with_remote
    a, b = "https://x.gov.au/a", "https://x.gov.au/b"
    remote, ours, out = tmp_path / "r.xlsx", tmp_path / "o.xlsx", tmp_path / "m.xlsx"
    _book(remote, [b], [a])          # remote archived a
    _book(ours, [a, b], [])          # ours still has a in Pending
    merge_xlsx_with_remote(remote, ours, out)
    ws = openpyxl.load_workbook(out)["Pending"]
    hdr = [c.value for c in ws[1]]
    urls = [row[hdr.index("URL")] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert urls == [b]


def test_row_merge_keeps_our_new_rows(tmp_path):
    import openpyxl
    from pipeline.xlsx_merge import merge_xlsx_with_remote
    a, b = "https://x.gov.au/a", "https://x.gov.au/b"
    remote, ours, out = tmp_path / "r.xlsx", tmp_path / "o.xlsx", tmp_path / "m.xlsx"
    _book(remote, [b], [])
    _book(ours, [a, b], [])
    merge_xlsx_with_remote(remote, ours, out)
    ws = openpyxl.load_workbook(out)["Pending"]
    hdr = [c.value for c in ws[1]]
    urls = sorted(row[hdr.index("URL")] for row in ws.iter_rows(min_row=2, values_only=True))
    assert urls == [a, b]


# ── 6. FSANZ: the Problem section, not the page chrome, names the hazard ───

_FSANZ_PAGE = """<html><head><title>Sunlife</title></head><body>
<h1>Sunlife Import &amp; Export Pty Ltd - Da Cha Gio Re Net Spring Roll Wrapper - 200g</h1>
<p>Published 30 September 2026</p>
<h2>Problem:</h2><p>The recall is due to the presence of an undeclared allergen (gluten).</p>
<h2>Food safety hazard:</h2><p>Any consumers who have a gluten allergy may have a reaction.</p>
<aside><h3>Other recalls</h3><p>Smoked salmon - Listeria monocytogenes</p></aside>
</body></html>"""


def test_fsanz_does_not_take_a_pathogen_from_the_page_chrome(monkeypatch):
    import scrapers.oceania.fsanz as mod

    class Resp:
        text = _FSANZ_PAGE

    monkeypatch.setattr(mod, "fetch", lambda session, url, *a, **k: Resp())
    s = mod.FSANZScraper.__new__(mod.FSANZScraper)
    s.session = None
    stats = {k: 0 for k in ("fetch_failed", "no_hazard", "no_date", "stale",
                            "no_title")}
    cutoff = dt.datetime(2026, 9, 1, tzinfo=dt.timezone.utc)
    rec = s._parse_recall_page("https://www.foodstandards.gov.au/food-recalls/"
                               "recall-alert/sunlife", cutoff, stats)
    assert rec is None and stats["no_hazard"] == 1


# ── 7. Dashboard Product filter uses the Alerts vocabulary ─────────────────

def test_dashboard_product_filter_matches_the_alert_vocabulary():
    """Operator 2026-10-02: a Product drop list on the dashboard, as Alerts
    has. Same categories and keywords (tools/alert_vocab.py), generated."""
    import subprocess
    html = (ROOT / "docs" / "index.html").read_text(encoding="utf-8")
    assert '<select id="f-prod"><option value="">All Products</option></select>' in html
    assert "if(prod&&!(r.prodCats||[]).includes(prod))return false;" in html
    assert "const PRODUCT_VOCAB = {" in html
    rc = subprocess.run([sys.executable, "tools/gen_alert_vocab.py", "--check"],
                        cwd=ROOT, capture_output=True, text=True)
    assert rc.returncode == 0, rc.stdout + rc.stderr
