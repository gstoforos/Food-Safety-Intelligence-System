# -*- coding: utf-8 -*-
"""tools/add_manual_row.py — the operator's hand-add path.

A hand-added row is the one row nobody can audit later by re-running a
scraper, so the tool has to be strict about two things: it must go into
Pending through the pipeline's own ``append_to_pending`` (same dedup,
same review), and it must never touch Recalls.
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from tools.add_manual_row import build_recall  # noqa: E402

TOOL = ROOT / "tools" / "add_manual_row.py"
# The worked example for tools/add_manual_row.py, kept as a TEST FIXTURE.
#
# It lived in a top-level rows/ directory for two days, which was a bad
# call on my part: nothing in the pipeline, the workflows or the site ever
# read it, so a reader finding rows/ at the repo root had no way to tell
# whether it was an input the system depended on. It is not. It is the
# hand-add spec for the 2026-08-28 Łowicz pesto row — already promoted,
# already in Recalls — retained so these tests have something real to
# exercise and so the row's provenance is reproducible.
ROW_JSON = ROOT / "tests" / "fixtures" / "gis-pesto-2026-08-28.json"


# --------------------------------------------------------------------------
# the shipped row
# --------------------------------------------------------------------------

@pytest.fixture(scope="module")
def pesto():
    return json.loads(ROW_JSON.read_text(encoding="utf-8"))


def test_the_pesto_row_file_is_present_and_valid_json(pesto):
    assert pesto["Source"] == "GIS (PL)"
    assert pesto["Date"] == "2026-08-28"


def test_pesto_row_is_tier_1_botulism_and_outbreak_flagged(pesto):
    r = build_recall(pesto)
    assert r.Pathogen == "Clostridium botulinum"
    assert int(r.Tier) == 1
    assert int(r.Outbreak) == 1
    assert r.Country == "Poland"
    assert r.Region == "Europe"


def test_pesto_row_points_at_the_regulator_not_the_news_story(pesto):
    """The news article is corroboration. The source of record is GIS."""
    assert pesto["URL"].startswith("https://www.gov.pl/web/gis/")
    assert "foodsafetynews" not in pesto["URL"]


def test_pesto_row_carries_the_batch_identifiers(pesto):
    """A botulism warning without the lot is not actionable in a store room."""
    prod = pesto["Product"]
    assert "L 605192FE I" in prod
    assert "180 g" in prod
    assert "08.2027" in prod


def test_pesto_notes_record_who_verified_what_and_what_is_unverified(pesto):
    notes = pesto["Notes"]
    assert "gov.pl" in notes or "GIS warning page" in notes
    # The Netherlands distribution comes from Food Safety News quoting
    # RASFF; the GIS page does not mention export and the RASFF
    # notification was not located. It must be marked, not asserted.
    assert "UNVERIFIED" in notes
    assert "Netherlands" in notes


def test_pesto_row_clears_the_publish_gate(pesto):
    from pipeline import _publish_gate as pg
    r = build_recall(pesto)
    row = {f: getattr(r, f) for f in
           ("Date", "Source", "Company", "Brand", "Product", "Pathogen",
            "Reason", "Class", "Country", "Region", "Tier", "Outbreak",
            "URL", "Notes")}
    assert pg.publish_blockers(row) == []


def test_pesto_pathogen_is_in_tier_1_scope(pesto):
    from pipeline import _pathogen_scope as ps
    r = build_recall(pesto)
    assert ps.is_in_scope(r.Pathogen)
    assert ps.is_always_tier1(r.Pathogen)


# --------------------------------------------------------------------------
# the tool's contract
# --------------------------------------------------------------------------

@pytest.mark.parametrize("field", ["Date", "Source", "Product", "URL"])
def test_missing_required_field_is_refused(pesto, field):
    spec = dict(pesto)
    spec.pop(field)
    with pytest.raises(SystemExit):
        build_recall(spec)


def test_a_bad_date_is_refused_rather_than_coerced(pesto):
    spec = dict(pesto, Date="28.08.2026")
    with pytest.raises(SystemExit):
        build_recall(spec)


def test_tier_is_computed_not_taken_from_the_json(pesto):
    """A hand-set Tier is exactly the kind of quiet error this avoids."""
    spec = dict(pesto, Tier=3)
    assert int(build_recall(spec).Tier) == 1


def test_pathogen_is_canonicalised(pesto):
    spec = dict(pesto, Pathogen="C. botulinum")
    assert build_recall(spec).Pathogen == "Clostridium botulinum"


def test_the_tool_never_writes_to_recalls():
    """Read the source: Recalls is loaded and passed straight back."""
    src = TOOL.read_text(encoding="utf-8")
    assert "append_to_pending" in src
    assert "promote_approved" not in src
    assert "save_xlsx(" not in src          # the Recalls-only writer


# --------------------------------------------------------------------------
# end to end, against a copy of the real workbook
# --------------------------------------------------------------------------

@pytest.fixture
def workbook_copy(tmp_path):
    src = ROOT / "docs" / "data" / "recalls.xlsx"
    if not src.exists():
        pytest.skip("recalls.xlsx not present")
    dst = tmp_path / "recalls.xlsx"
    dst.write_bytes(src.read_bytes())
    return dst


def _counts(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    n = {s: max(0, len(list(wb[s].iter_rows(values_only=True))) - 1)
         for s in ("Recalls", "Pending")}
    wb.close()
    return n


def _run(*argv):
    return subprocess.run([sys.executable, str(TOOL), *argv],
                          cwd=ROOT, capture_output=True, text=True)


@pytest.fixture
def novel_row(tmp_path, pesto):
    """The pesto row with a URL the register has never seen.

    AUDIT 2026-09-16 — the end-to-end tests used to queue the shipped pesto
    row itself. That worked only until the row was promoted; now that it is
    in Recalls, append_to_pending correctly SKIPS it and the tests failed
    for the one reason that means the tool is working.

    So the pipeline behaviour is split in two, and both halves are asserted:
    a genuinely new row queues (below), and the shipped row is refused
    because it is already approved (test_shipped_row_is_now_a_no_op).
    """
    spec = dict(pesto)
    # A different URL is NOT enough. append_to_pending also runs a
    # near-duplicate index on (company, date, pathogen) — the guard added
    # 2026-04-29 against hallucinated gap-finder URLs — so a re-skinned
    # pesto row is still recognised as the approved one. Change the
    # identity, not just the link.
    spec["URL"] = "https://www.gov.pl/web/gis/unit-test-only-never-real"
    spec["Company"] = "UNIT TEST Sp. z o.o. (not a real firm)"
    spec["Date"] = "2026-08-27"
    spec["Product"] = "UNIT TEST pesto — synthetic row, never publish"
    f = tmp_path / "novel.json"
    f.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    return str(f)


def test_shipped_row_is_now_a_no_op(workbook_copy):
    """The pesto row is in Recalls. Re-running must add nothing.

    This is append_to_pending's own dedup rule — already approved, skip
    silently — and it is what stops a hand-add being queued twice by an
    operator who does not remember running it.
    """
    before = _counts(workbook_copy)
    r = _run(str(ROW_JSON), "--xlsx", str(workbook_copy))
    assert r.returncode == 0, r.stderr
    assert "Nothing queued" in r.stdout
    assert _counts(workbook_copy) == before


def test_dry_run_writes_nothing(workbook_copy, novel_row):
    before = workbook_copy.read_bytes()
    r = _run(novel_row, "--xlsx", str(workbook_copy), "--dry-run")
    assert r.returncode == 0, r.stderr
    assert "DRY RUN" in r.stdout
    assert workbook_copy.read_bytes() == before


def test_queues_one_row_into_pending_and_leaves_recalls_alone(workbook_copy, novel_row):
    before = _counts(workbook_copy)
    r = _run(novel_row, "--xlsx", str(workbook_copy))
    assert r.returncode == 0, r.stderr
    after = _counts(workbook_copy)
    assert after["Pending"] == before["Pending"] + 1
    assert after["Recalls"] == before["Recalls"]


def test_the_queued_row_is_status_pending(workbook_copy, novel_row):
    _run(novel_row, "--xlsx", str(workbook_copy))
    wb = load_workbook(workbook_copy, read_only=True, data_only=True)
    rows = list(wb["Pending"].iter_rows(values_only=True))
    hdr = [str(c or "") for c in rows[0]]
    hits = [dict(zip(hdr, x)) for x in rows[1:]
            if "pesto" in str(dict(zip(hdr, x)).get("Product", "")).lower()]
    wb.close()
    assert len(hits) == 1
    assert str(hits[0]["Status"]).lower() == "pending"
    assert int(hits[0]["Tier"]) == 1


def test_running_twice_does_not_duplicate(workbook_copy, novel_row):
    """append_to_pending's own rule: already pending -> skip."""
    _run(novel_row, "--xlsx", str(workbook_copy))
    mid = _counts(workbook_copy)
    r = _run(novel_row, "--xlsx", str(workbook_copy))
    assert r.returncode == 0, r.stderr
    assert _counts(workbook_copy)["Pending"] == mid["Pending"]
    assert "Nothing queued" in r.stdout
