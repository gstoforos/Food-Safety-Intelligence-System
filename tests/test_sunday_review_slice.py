# -*- coding: utf-8 -*-
"""The Sunday email must report the week that closed, not the one opening.

THE BUG (audit 2026-09-20)
--------------------------
Every Sunday manual-review email said the same thing:

    0 recalls added
    No new recalls promoted into the Recalls sheet this week.

It was not true. In the week to 2026-09-20 the Weekly_Review sheet held
**15 rows, 13 of them Tier 1**.

``review_day_for()`` answers "a row promoted right now appears in WHICH
email?" — the correct stamp to write at promotion time, and it rolls
over to next Sunday once the 17:00 Athens cutoff passes.

``export_week_slice()`` reused it as the FILTER for the email being
sent. The mailer fires at Sunday 17:00 — the exact moment of the
rollover — so the export looked in next week's bucket, which is empty by
construction.

Measured: the 15 rows all carry ``Week_Added = 2026-09-20``. At 17:14
Athens ``review_day_for()`` returns **2026-09-27**. Zero matches.

It is also why the live file read ``week_end: 2026-09-13`` with
``generated_utc: 2026-09-06`` — the same off-by-one-week, a fortnight
earlier, never regenerated since.

The freshness banner in the email ("JSON is 335.5h old") was right and
firing the whole time. The number behind it was the one nobody could
explain.
"""

from __future__ import annotations

import sys
from datetime import datetime, timezone
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from pipeline.weekly_review_capture import (        # noqa: E402
    review_day_for,
    review_day_just_closed,
)

#: 2026-09-20 is a Sunday. Athens is UTC+3 in September.
SUN_1714 = datetime.fromisoformat("2026-09-20T14:14:00+00:00")   # 17:14 Athens
SUN_1000 = datetime.fromisoformat("2026-09-20T07:00:00+00:00")   # 10:00 Athens
MON_0900 = datetime.fromisoformat("2026-09-21T06:00:00+00:00")
WED_1200 = datetime.fromisoformat("2026-09-23T09:00:00+00:00")
SAT_2300 = datetime.fromisoformat("2026-09-19T20:00:00+00:00")   # 23:00 Athens


# --------------------------------------------------------------------------
# the incident
# --------------------------------------------------------------------------

def test_the_mailer_moment_reports_the_week_that_just_closed():
    """THE regression. At Sunday 17:00+ the email is about TODAY's bucket."""
    assert review_day_just_closed(SUN_1714).isoformat() == "2026-09-20"


def test_and_the_stamping_function_has_already_rolled_over_by_then():
    """Both are correct; they answer different questions. This is the trap."""
    assert review_day_for(SUN_1714).isoformat() == "2026-09-27"
    assert review_day_for(SUN_1714) != review_day_just_closed(SUN_1714), (
        "if these ever agree at the mailer moment, the distinction has been "
        "collapsed and the Sunday email will go empty again")


def test_export_defaults_to_the_closed_week_not_the_open_one():
    """Guard the wiring, not just the helper."""
    import inspect
    from pipeline import weekly_review_capture as wrc
    src = inspect.getsource(wrc.export_week_slice)
    assert "review_day_just_closed()" in src, (
        "export_week_slice must default to the CLOSED week; defaulting to "
        "review_day_for() is what produced '0 recalls added' every week")


# --------------------------------------------------------------------------
# the boundary, from both sides
# --------------------------------------------------------------------------

@pytest.mark.parametrize("when,expect", [
    (SUN_1714, "2026-09-20"),   # after the cutoff: this week has closed
    (SUN_1000, "2026-09-13"),   # before it: still filling, last closed is prior
    (MON_0900, "2026-09-20"),
    (WED_1200, "2026-09-20"),
    (SAT_2300, "2026-09-13"),   # Saturday night: Sunday has not happened yet
])
def test_closed_week_across_the_cycle(when, expect):
    assert review_day_just_closed(when).isoformat() == expect


def test_it_always_lands_on_a_sunday():
    from datetime import timedelta
    base = datetime(2026, 9, 1, tzinfo=timezone.utc)
    for h in range(0, 24 * 21, 5):          # three weeks, every five hours
        d = review_day_just_closed(base + timedelta(hours=h))
        assert d.weekday() == 6, "%s is a %s, not a Sunday" % (d, d.strftime("%A"))


def test_the_closed_week_is_never_in_the_future():
    from datetime import timedelta
    base = datetime(2026, 9, 1, tzinfo=timezone.utc)
    for h in range(0, 24 * 21, 7):
        now = base + timedelta(hours=h)
        assert review_day_just_closed(now) <= now.date(), (
            "the email would be reporting on a week that has not happened")


# --------------------------------------------------------------------------
# end to end against the real sheet
# --------------------------------------------------------------------------

def test_the_slice_finds_the_rows_that_are_actually_there():
    """The sheet had 15 rows and the email said 0. Never again."""
    from openpyxl import load_workbook
    xlsx = ROOT / "docs" / "data" / "recalls.xlsx"
    if not xlsx.exists():
        pytest.skip("recalls.xlsx not present")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if "Weekly_Review" not in wb.sheetnames:
        wb.close()
        pytest.skip("no Weekly_Review sheet")
    raw = list(wb["Weekly_Review"].iter_rows(values_only=True))
    hdr = [str(c or "") for c in raw[0]]
    wb.close()
    if "Week_Added" not in hdr or len(raw) < 2:
        pytest.skip("sheet empty or has no Week_Added column")

    idx = hdr.index("Week_Added")
    stamps = {str(r[idx] or "") for r in raw[1:] if r[idx]}
    assert stamps, "Weekly_Review rows carry no Week_Added stamp"

    from pipeline.weekly_review_capture import export_week_slice
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as fh:
        out = Path(fh.name)
    payload = export_week_slice(xlsx_path=xlsx, json_path=out)
    out.unlink(missing_ok=True)

    if payload["week_end"] in stamps:
        assert payload["row_count"] > 0, (
            "the slice for %s matched 0 rows, but the sheet has rows stamped "
            "with exactly that week — this is the 2026-09-20 bug"
            % payload["week_end"])
