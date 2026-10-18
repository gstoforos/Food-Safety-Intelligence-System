"""The Recalls row-union must not put back a row the remote side removed
(2026-10-18).

On 2026-10-17 an operator upload removed two NAFDAC rows and archived two
CFS (HK) duplicates; the next writer, holding the pre-upload workbook, lost
the push race and the union appended all four back. An OUR-side row is now
dropped only when it is absent from the remote's Recalls AND its URL is on
record in the remote's Rejected or Weekly_Rejected.
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from pipeline.xlsx_merge import merge_xlsx_with_remote
from tests.conftest import (NEWS_COLS, PENDING_COLS, RECALLS_COLS,
                            WEEKLY_REJECTED_COLS)

pytestmark = pytest.mark.slow

REMOVED = [
    "https://nafdac.gov.ng/public-alert-no-023-2025-recall-of-multivitamin-supplements/",
    "https://nafdac.gov.ng/public-alert-no-027-2022-adulterated-lots-of-japata-bitters/",
    "https://www.cfs.gov.hk/english/whatsnew/whatsnew_fa/2026_633.html",
    "https://www.cfs.gov.hk/english/whatsnew/whatsnew_fa/2026_632.html",
]
KEEP = "https://www.cfs.gov.hk/english/press/20261012_12665.html"


def _urls(path, sheet="Recalls"):
    ws = load_workbook(path)[sheet]
    hdr = [c.value for c in ws[1]]
    i = hdr.index("URL")
    return {r[i] for r in ws.iter_rows(min_row=2, values_only=True) if r[i]}


def _row(base, url):
    return {**base, "URL": url}


def _books(factory, base, ours_urls, remote_urls, archived, sheet="Weekly_Rejected"):
    wj = [_row(base, u) for u in archived]
    remote = factory("remote.xlsx", {
        "Recalls": (RECALLS_COLS, [_row(base, u) for u in remote_urls]),
        "Pending": (PENDING_COLS, []),
        "NEWS": (NEWS_COLS, []),
        sheet: (WEEKLY_REJECTED_COLS if sheet == "Weekly_Rejected" else RECALLS_COLS + ["RejectReason"], wj),
    })
    ours = factory("ours.xlsx", {
        "Recalls": (RECALLS_COLS, [_row(base, u) for u in ours_urls]),
        "Pending": (PENDING_COLS, []),
        "NEWS": (NEWS_COLS, []),
    })
    return remote, ours


@pytest.mark.parametrize("url", REMOVED)
def test_a_removed_and_archived_row_stays_out(url, tmp_xlsx_factory, tmp_path, sample_recall_row):
    remote, ours = _books(tmp_xlsx_factory, sample_recall_row,
                          ours_urls=[KEEP, url], remote_urls=[KEEP], archived=[url])
    out = tmp_path / "out.xlsx"
    merge_xlsx_with_remote(remote, ours, out)
    got = _urls(out)
    assert url not in got and KEEP in got


def test_the_rejected_sheet_counts_too(tmp_xlsx_factory, tmp_path, sample_recall_row):
    url = REMOVED[0]
    remote, ours = _books(tmp_xlsx_factory, sample_recall_row, ours_urls=[url],
                          remote_urls=[KEEP], archived=[url], sheet="Rejected")
    out = tmp_path / "out.xlsx"
    merge_xlsx_with_remote(remote, ours, out)
    assert url not in _urls(out)


def test_a_new_row_from_the_losing_run_still_merges(tmp_xlsx_factory, tmp_path, sample_recall_row):
    new = "https://www.fda.gov/safety/recalls-market-withdrawals-safety-alerts/new-recall"
    remote, ours = _books(tmp_xlsx_factory, sample_recall_row, ours_urls=[KEEP, new],
                          remote_urls=[KEEP], archived=[REMOVED[0]])
    out = tmp_path / "out.xlsx"
    merge_xlsx_with_remote(remote, ours, out)
    assert new in _urls(out)


def test_a_remote_row_is_never_dropped_even_if_archived(tmp_xlsx_factory, tmp_path, sample_recall_row):
    # published AND archived on the remote: a contradiction for the morning
    # pass to annotate, not for a merge to delete
    url = REMOVED[2]
    remote, ours = _books(tmp_xlsx_factory, sample_recall_row, ours_urls=[url],
                          remote_urls=[url], archived=[url])
    out = tmp_path / "out.xlsx"
    merge_xlsx_with_remote(remote, ours, out)
    assert url in _urls(out)
