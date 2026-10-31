# -*- coding: utf-8 -*-
"""A push retry must not throw away everything except the xlsx.

THE BUG THIS EXISTS TO PREVENT (morning fix 2026-10-30)
=======================================================
scripts/safe_push.sh loses the push race often — many data-writer
workflows share main. On rejection it did this::

    cp "$XLSX" /tmp/ours_recalls.xlsx
    git reset --hard HEAD~1
    git pull --rebase --autostash origin main
    python -m pipeline.xlsx_merge ...        # xlsx rows unioned back
    stage_existing
    git commit -m "$MSG (merge attempt $attempt)"

``git reset --hard`` reverts the WORKING TREE for every path in our
commit. Only the xlsx had been copied aside. So every other path the run
produced went back to the remote's version, ``stage_existing`` then found
nothing to stage for it, and the retry commit carried the xlsx ALONE.

MEASURED, on main, 2026-10-30: all 170 ``(merge attempt N)`` commits in
the preceding 30 days carried ``docs/data/recalls.xlsx`` and nothing
else. Not one carried a single jsonl file. The Nordic and Central EU gap
finders retry every day, so their per-country logs had stopped being
committed altogether.

The script's own header called this correct, on the grounds that the
extra paths are "regenerated each run". ``run_log.jsonl`` is not: it is
APPEND-ONLY, one record per country per exit, and it is what
tests/test_no_country_goes_dark.py and tools/fleet_shard.catchup_codes
read to decide whether a country has gone dark.

What that cost: the fleet ran ``us`` on 2026-10-29 in a shard of nine,
the run wrote its record, the retry discarded it, and
test_no_country_goes_dark went red with "'us' last ran 2026-10-14, 15
days ago" for a country that was running normally. catchup_codes reads
the same log, so ``us`` was force-added to the catch-up every day from
2026-10-24, burning one of three daily slots that genuinely stale
countries needed.

These tests drive the real script against two real clones of a real bare
origin, with the race actually lost, because the defect lived in the
interaction between reset, rebase and staging — not in anything a mocked
git would have shown.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

SCRIPT = ROOT / "scripts" / "safe_push.sh"
XLSX_REL = "docs/data/recalls.xlsx"
GAP_DIR = "docs/data/gap_finder_us"

pytestmark = pytest.mark.skipif(
    shutil.which("git") is None or shutil.which("awk") is None,
    reason="needs git and awk on PATH",
)

SHEETS = ("Recalls", "Pending", "Weekly_Review", "Weekly_Rejected",
          "Rejected")
HEADER = ["Date", "Source", "Company", "Product", "URL"]


def _git(cwd: Path, *args: str) -> str:
    env = dict(os.environ)
    env.update(
        GIT_AUTHOR_NAME="t", GIT_AUTHOR_EMAIL="t@t.t",
        GIT_COMMITTER_NAME="t", GIT_COMMITTER_EMAIL="t@t.t",
        GIT_CONFIG_GLOBAL=str(cwd / ".gitconfig-none"),
        GIT_CONFIG_SYSTEM=str(cwd / ".gitconfig-none"),
    )
    out = subprocess.run(("git",) + args, cwd=str(cwd), env=env, text=True,
                         capture_output=True)
    assert out.returncode == 0, f"git {args}: {out.stdout}\n{out.stderr}"
    return out.stdout


def _write_xlsx(path: Path, urls) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEETS[0]
    ws.append(HEADER)
    for u in urls:
        ws.append(["2026-10-30", "S", "C", "P", u])
    for name in SHEETS[1:]:
        wb.create_sheet(name).append(HEADER)
    wb.save(path)


def _xlsx_urls(path: Path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Recalls"].values)
    return [r[HEADER.index("URL")] for r in rows[1:]]


def _add_row(path: Path, url: str) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(path)
    wb["Recalls"].append(["2026-10-30", "S", "C", "P", url])
    wb.save(path)


@pytest.fixture()
def race(tmp_path):
    """A bare origin and two clones, with pipeline/ importable in each."""
    origin = tmp_path / "origin.git"
    _git(tmp_path, "init", "-q", "--bare", "--initial-branch=main",
         str(origin))

    seed = tmp_path / "seed"
    _git(tmp_path, "clone", "-q", str(origin), str(seed))
    (seed / GAP_DIR).mkdir(parents=True)
    (seed / "scripts").mkdir()
    shutil.copy(SCRIPT, seed / "scripts" / "safe_push.sh")
    shutil.copytree(ROOT / "pipeline", seed / "pipeline",
                    ignore=shutil.ignore_patterns("__pycache__"))
    _write_xlsx(seed / XLSX_REL, ["https://e/1"])
    (seed / GAP_DIR / "run_log.jsonl").write_text(
        json.dumps({"country_code": "us",
                    "finished_at": "2026-10-14T10:53:40+00:00"}) + "\n",
        encoding="utf-8")
    (seed / GAP_DIR / "candidates.jsonl").write_text(
        '{"cand": "old"}\n', encoding="utf-8")
    (seed / GAP_DIR / "doomed.jsonl").write_text(
        '{"stale": true}\n', encoding="utf-8")
    _git(seed, "add", "-A")
    _git(seed, "commit", "-qm", "base")
    _git(seed, "push", "-q", "origin", "main")

    winner = tmp_path / "winner"
    loser = tmp_path / "loser"
    _git(tmp_path, "clone", "-q", str(origin), str(winner))
    _git(tmp_path, "clone", "-q", str(origin), str(loser))
    return winner, loser


def _winner_pushes(winner: Path, *, also_logs: bool) -> None:
    """The concurrent writer that gets to main first."""
    _add_row(winner / XLSX_REL, "https://e/2")
    if also_logs:
        with (winner / GAP_DIR / "run_log.jsonl").open(
                "a", encoding="utf-8") as fh:
            fh.write(json.dumps({"country_code": "us",
                                 "finished_at": "2026-10-30T18:00:00+00:00",
                                 "who": "concurrent-writer"}) + "\n")
    _git(winner, "add", "-A")
    _git(winner, "commit", "-qm", "other writer")
    _git(winner, "push", "-q", "origin", "main")


def _loser_runs_safe_push(loser: Path):
    """Our run: a new xlsx row, a new log record, regenerated jsonl,
    and one file the run deleted."""
    _add_row(loser / XLSX_REL, "https://e/3")
    with (loser / GAP_DIR / "run_log.jsonl").open(
            "a", encoding="utf-8") as fh:
        fh.write(json.dumps({"country_code": "us",
                             "finished_at": "2026-10-30T10:53:00+00:00",
                             "status": "completed"}) + "\n")
    (loser / GAP_DIR / "candidates.jsonl").write_text(
        '{"cand": "fresh"}\n', encoding="utf-8")
    (loser / GAP_DIR / "doomed.jsonl").unlink()

    env = dict(os.environ)
    env.update(PYTHONPATH=str(loser), HOME=str(loser))
    out = subprocess.run(
        ["bash", "scripts/safe_push.sh",
         "Gap finder fleet: 2026-10-30 shard of 1 countries",
         XLSX_REL, GAP_DIR + "/"],
        cwd=str(loser), env=env, text=True, capture_output=True)
    assert out.returncode == 0, (
        f"safe_push.sh exited {out.returncode}\n"
        f"--- stdout ---\n{out.stdout}\n--- stderr ---\n{out.stderr}")
    return out


def _committed_paths(repo: Path):
    return set(_git(repo, "show", "--name-only", "--pretty=format:",
                    "HEAD").split())


def test_the_retry_path_is_actually_exercised(race):
    """Guards the test itself: if the push stopped being rejected the
    assertions below would pass vacuously on a first-attempt commit."""
    winner, loser = race
    _winner_pushes(winner, also_logs=False)
    out = _loser_runs_safe_push(loser)
    assert "rejected — syncing remote and row-merging" in out.stdout
    assert "(merge attempt 1)" in _git(loser, "log", "-1", "--format=%s")


def test_a_retry_commit_carries_the_non_xlsx_paths_too(race):
    """The 170-commits-in-30-days bug, in one assertion."""
    winner, loser = race
    _winner_pushes(winner, also_logs=False)
    _loser_runs_safe_push(loser)

    paths = _committed_paths(loser)
    assert f"{GAP_DIR}/run_log.jsonl" in paths, (
        f"the retry commit carried only {sorted(paths)} — the run's "
        f"append-only run_log was reverted by `git reset --hard` and the "
        f"country now looks dark to tests/test_no_country_goes_dark.py "
        f"and to tools/fleet_shard.catchup_codes")
    assert f"{GAP_DIR}/candidates.jsonl" in paths
    assert paths != {XLSX_REL}, (
        "the retry commit is xlsx-only again — this is the exact shape of "
        "every merge-attempt commit on main for the 30 days before "
        "2026-10-30")


def test_the_run_log_record_survives_the_retry(race):
    winner, loser = race
    _winner_pushes(winner, also_logs=False)
    _loser_runs_safe_push(loser)

    lines = [json.loads(l) for l in
             (loser / GAP_DIR / "run_log.jsonl").read_text(
                 encoding="utf-8").splitlines() if l.strip()]
    stamps = [r.get("finished_at") for r in lines]
    assert "2026-10-30T10:53:00+00:00" in stamps, (
        f"our run's record is gone after the retry; log holds {stamps}")


def test_a_concurrent_writers_log_record_is_not_overwritten(race):
    """Ours must win without taking the remote's records with it.

    Both sides append to the same log. Keeping only ours would be the
    same defect with the sign flipped, so the restore unions the lines.
    """
    winner, loser = race
    _winner_pushes(winner, also_logs=True)
    _loser_runs_safe_push(loser)

    stamps = [json.loads(l).get("finished_at") for l in
              (loser / GAP_DIR / "run_log.jsonl").read_text(
                  encoding="utf-8").splitlines() if l.strip()]
    assert "2026-10-30T10:53:00+00:00" in stamps, "our record was lost"
    assert "2026-10-30T18:00:00+00:00" in stamps, (
        f"the concurrent writer's record was overwritten by ours; log "
        f"holds {stamps}")
    assert len(stamps) == len(set(stamps)), (
        f"the union duplicated records: {stamps}")


def test_regenerated_jsonl_keeps_this_runs_version(race):
    """candidates.jsonl IS a regenerated snapshot — ours is the fresh
    one, so it must not come back as the remote's stale copy."""
    winner, loser = race
    _winner_pushes(winner, also_logs=False)
    _loser_runs_safe_push(loser)
    assert "fresh" in (loser / GAP_DIR / "candidates.jsonl").read_text(
        encoding="utf-8")


def test_a_deletion_is_not_resurrected_by_the_retry(race):
    winner, loser = race
    _winner_pushes(winner, also_logs=False)
    _loser_runs_safe_push(loser)
    assert not (loser / GAP_DIR / "doomed.jsonl").exists(), (
        "the rebase brought back a file this run deleted")


def test_the_xlsx_row_union_still_works(race):
    """The behaviour that was already right must stay right."""
    winner, loser = race
    _winner_pushes(winner, also_logs=False)
    _loser_runs_safe_push(loser)
    urls = _xlsx_urls(loser / XLSX_REL)
    assert "https://e/2" in urls, "the winner's row was dropped"
    assert "https://e/3" in urls, "our row was dropped"


def test_the_script_still_copies_the_xlsx_aside(race):
    """A guard on the fix, not on git: preserve/restore must not have
    replaced the xlsx copy-aside, which is what makes the row-union
    possible."""
    body = SCRIPT.read_text(encoding="utf-8")
    assert "cp \"$XLSX\" /tmp/ours_recalls.xlsx" in body
    assert "preserve_our_non_xlsx_work" in body
    assert "restore_our_non_xlsx_work" in body
    # Order matters, and only CODE lines carry it — the prose above the
    # retry block names `git reset --hard HEAD~1` too.
    code = [l.strip() for l in body.splitlines()
            if l.strip() and not l.lstrip().startswith("#")]
    reset = code.index("git reset --hard HEAD~1")
    preserve = code.index("preserve_our_non_xlsx_work")
    restore = code.index("restore_our_non_xlsx_work")
    assert preserve < reset, (
        "preserve must run while our commit is still HEAD, otherwise there "
        "is nothing left to preserve")
    assert restore > reset, "restore must run after the pull"
