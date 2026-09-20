# -*- coding: utf-8 -*-
"""The register must never lose published rows without an explicit decision.

THE INCIDENT (2026-09-20)
-------------------------
Commit ``02d97dde`` — message "Add files via upload" — took Recalls from
**1751 to 1722**. Twenty-nine published rows disappeared, **nineteen of
them Tier 1**, covering every promotion made on 2026-09-18 and 09-19.

Nothing failed. No workflow went red, no email was sent, no test
complained. The register simply got smaller, the dashboard showed the
newest recall as 17 September, and the week looked stalled at 25 rows.
The pipeline upstream had been working perfectly the whole time.

The cause was a stale data file: ``docs/data/recalls.xlsx`` had been
shipped inside a dated zip, and the zip was uploaded two days later,
overwriting the live register with a snapshot of itself. Confirmed by a
byte-for-byte comparison.

WHY A TEST AND NOT A CONVENTION
-------------------------------
"Never ship the register in a zip" is a rule someone has to remember
every single time. This is the check that does not.

Deletions are legitimate — a duplicate batch, a mis-scraped row, a
rejected set. What is never legitimate is a deletion nobody declared. So
the count may fall, but only when the commit that lowers it says so.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

XLSX = ROOT / "docs" / "data" / "recalls.xlsx"

#: A commit may remove rows when its message says so in as many words.
DELETION_MARKERS = (
    "remove-rows", "delete-rows", "reject-batch",
    "deliberate row removal", "approved deletion",
)


def _rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    raw = list(wb["Recalls"].iter_rows(values_only=True))
    hdr = [str(c or "") for c in raw[0]]
    out = [dict(zip(hdr, r)) for r in raw[1:]]
    wb.close()
    return out


def _git(*args):
    return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True)


def _previous_workbook():
    if _git("rev-parse", "HEAD~1").returncode:
        return None
    show = subprocess.run(["git", "show", "HEAD~1:docs/data/recalls.xlsx"],
                          cwd=ROOT, capture_output=True)
    if show.returncode or not show.stdout:
        return None
    fh = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    fh.write(show.stdout)
    fh.close()
    return Path(fh.name)


def test_the_workbook_is_readable():
    assert XLSX.exists(), "docs/data/recalls.xlsx is missing"
    assert len(_rows(XLSX)) > 0


def test_recalls_did_not_shrink_without_saying_so():
    if not (ROOT / ".git").exists():
        pytest.skip("not a git checkout")
    prev = _previous_workbook()
    if prev is None:
        pytest.skip("no previous revision of the workbook")

    before, after = _rows(prev), _rows(XLSX)
    Path(prev).unlink(missing_ok=True)

    keys = {str(r.get("URL") or "").strip().lower() for r in after}
    gone = [r for r in before
            if str(r.get("URL") or "").strip().lower() not in keys]
    if not gone:
        return

    msg = _git("log", "-1", "--format=%B").stdout.lower()
    if any(m in msg for m in DELETION_MARKERS):
        return

    tier1 = sum(1 for r in gone if str(r.get("Tier")).strip() == "1")
    sample = "\n  ".join(
        "%s | %s | %s" % (str(r.get("Date"))[:10], str(r.get("Source"))[:16],
                          str(r.get("Product"))[:54])
        for r in gone[:8])
    raise AssertionError(
        "%d published row(s) vanished from Recalls (%d Tier 1) and this "
        "commit's message does not declare a deletion.\n\nThis is the "
        "2026-09-20 shape: an 'Add files via upload' commit overwrote the "
        "live register with an older snapshot and removed 29 rows, 19 Tier 1, "
        "in silence.\n\nIf intended, say so in the commit message using one "
        "of: %s\n\nRemoved (first 8 of %d):\n  %s"
        % (len(gone), tier1, ", ".join(DELETION_MARKERS), len(gone), sample))


def test_the_json_mirror_matches_the_workbook():
    import json
    js = ROOT / "docs" / "data" / "recalls.json"
    if not js.exists():
        pytest.skip("recalls.json not present")
    n_json = len(json.loads(js.read_text(encoding="utf-8-sig")))
    n_xlsx = len(_rows(XLSX))
    assert n_json == n_xlsx, (
        "recalls.json has %d rows, recalls.xlsx has %d — the site is serving "
        "a different register from the workbook" % (n_json, n_xlsx))
